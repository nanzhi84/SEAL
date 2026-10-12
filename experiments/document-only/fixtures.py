"""Independent synthetic resources for the unified document user journey."""

import json
import struct
import zlib
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

ROOT = Path(__file__).resolve().parents[2]


def png_chunk(kind, body):
    return struct.pack(">I", len(body)) + kind + body + struct.pack(">I", zlib.crc32(kind + body))


PNG = (
    b"\x89PNG\r\n\x1a\n"
    + png_chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
    + png_chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00"))
    + png_chunk(b"IEND", b"")
)
JSON_VALUE = {"year": 2026, "authority": None, "tags": ["policy", "science"]}


def xlsx_bytes():
    """Author a minimal two-sheet workbook without using the runtime reader."""
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    relationship = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package = "http://schemas.openxmlformats.org/package/2006/relationships"
    files = {
        "[Content_Types].xml": (
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            '<Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            "</Types>"
        ),
        "_rels/.rels": (
            f'<Relationships xmlns="{package}"><Relationship Id="rId1" '
            f'Type="{relationship}/officeDocument" Target="xl/workbook.xml"/></Relationships>'
        ),
        "xl/workbook.xml": (
            f'<workbook xmlns="{main}" xmlns:r="{relationship}"><sheets>'
            '<sheet name="First" sheetId="1" r:id="rId1"/>'
            '<sheet name="Second" sheetId="2" r:id="rId2"/></sheets></workbook>'
        ),
        "xl/_rels/workbook.xml.rels": (
            f'<Relationships xmlns="{package}">'
            f'<Relationship Id="rId1" Type="{relationship}/worksheet" Target="worksheets/sheet1.xml"/>'
            f'<Relationship Id="rId2" Type="{relationship}/worksheet" Target="worksheets/sheet2.xml"/>'
            "</Relationships>"
        ),
        "xl/worksheets/sheet1.xml": (
            f'<worksheet xmlns="{main}"><sheetData><row r="1">'
            '<c r="A1" t="inlineStr"><is><t>name</t></is></c>'
            '<c r="B1" t="inlineStr"><is><t>amount</t></is></c></row>'
            '<row r="2"><c r="A2" t="inlineStr"><is><t>Public entity</t></is></c>'
            '<c r="B2"><v>42</v></c></row></sheetData></worksheet>'
        ),
        "xl/worksheets/sheet2.xml": (
            f'<worksheet xmlns="{main}"><sheetData><row r="1">'
            '<c r="A1" t="inlineStr"><is><t>Merged</t></is></c></row>'
            '<row r="2"><c r="A2" t="b"><v>1</v></c></row></sheetData>'
            '<mergeCells count="1"><mergeCell ref="A1:B1"/></mergeCells></worksheet>'
        ),
    }
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as package_file:
        for path, value in files.items():
            package_file.writestr(path, value)
    return stream.getvalue()


def resources(pdf, docx, xls):
    html = (
        "<html><h1>Unified public notice</h1><article><p>Before the resources.</p>"
        '<img src="/documents/image.png" alt="One public pixel">'
        "<table><caption>Public table</caption><tr><th>Name</th><th>Amount</th></tr>"
        "<tr><td>Public entity</td><td>42</td></tr></table>"
        '<a class="attachment" href="/documents/shared.xlsx">Shared workbook</a>'
        "<p>After the resources.</p></article></html>"
    ).encode()
    return {
        "/robots.txt": ("text/plain", b"User-agent: *\nAllow: /\n"),
        "/documents/list": (
            "text/html; charset=utf-8",
            b'<html><nav class="documents"><a href="/documents/notice">Notice</a>'
            b'<a href="/documents/also">Second reference</a></nav>'
            b'<a class="next" href="/documents/page2">Next</a></html>',
        ),
        "/documents/page2": (
            "text/html; charset=utf-8",
            b'<html><nav class="documents">'
            + b"".join(
                f'<a href="/documents/{name}">{name}</a>'.encode()
                for name in (
                    "report.pdf",
                    "report.docx",
                    "workbook.xls",
                    "public.csv",
                    "plain.txt",
                    "data.json",
                    "array.json",
                    "scalar.json",
                    "null.json",
                    "image.png",
                )
            )
            + b"</nav></html>",
        ),
        "/documents/notice": ("text/html; charset=utf-8", html),
        "/documents/also": (
            "text/html; charset=utf-8",
            b"<h1>Another referring page</h1><article>Shared input"
            b'<a class="attachment" href="/documents/shared.xlsx">The same workbook</a></article>',
        ),
        "/documents/shared.xlsx": (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            xlsx_bytes(),
        ),
        "/documents/report.pdf": ("application/pdf", pdf),
        "/documents/report.docx": (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            docx,
        ),
        "/documents/workbook.xls": ("application/vnd.ms-excel", xls),
        "/documents/public.csv": (
            "text/csv; charset=utf-8",
            b'name,amount,comment\n"Public, entity",42,"quoted ""value"""\nEmpty,,\n',
        ),
        "/documents/plain.txt": ("text/plain; charset=utf-8", b"Plain notice\nPublic plain body"),
        "/documents/data.json": (
            "application/json; charset=utf-8",
            json.dumps(JSON_VALUE).encode(),
        ),
        "/documents/array.json": ("application/json", b'[1,{"kind":"public"},null]'),
        "/documents/related.json": (
            "application/json",
            b'{"year":2027,"authority":"A second archived resource"}',
        ),
        "/documents/scalar.json": ("application/json", b'"Public scalar"'),
        "/documents/null.json": ("application/json", b"null"),
        "/documents/image.png": ("image/png", PNG),
        "/documents/bad.pdf": ("application/pdf", b"damaged synthetic PDF"),
        "/documents/bad.docx": (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            b"damaged synthetic DOCX",
        ),
        "/documents/bad.xlsx": (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            b"damaged synthetic XLSX",
        ),
        "/documents/bad.json": ("application/json", b'{"invalid":'),
        "/documents/unsupported.bin": ("application/octet-stream", b"unsupported binary bytes"),
        "/documents/ref-parent": (
            "text/html; charset=utf-8",
            b'<article><h1>References</h1><a class="attachment" href="/documents/bad.pdf">'
            b'Broken PDF</a><a class="attachment" href="/documents/missing.txt">'
            b'Unavailable text</a><a class="attachment" href="/documents/not-fetched.txt">'
            b"Not scheduled</a></article>",
        ),
        "/documents/not-fetched.txt": ("text/plain", b"This resource must never be downloaded"),
        "/documents/access-denied": (
            "text/html; charset=utf-8",
            b'<title>Access denied</title><div class="documents">'
            b'<a href="/documents/plain.txt">Still discover the actual resource</a></div>',
        ),
    }


def pdf_bytes():
    # Fixed, text-layer PDF; the oracle is independent of pypdf extraction.
    stream = b"BT /F1 18 Tf 50 750 Td (Synthetic PDF notice) Tj ET"
    parts = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    data = b"%PDF-1.4\n"
    offsets = [0]
    for i, part in enumerate(parts, 1):
        offsets.append(len(data))
        data += f"{i} 0 obj\n".encode() + part + b"\nendobj\n"
    start = len(data)
    data += b"xref\n0 6\n0000000000 65535 f \n"
    data += b"".join(f"{n:010d} 00000 n \n".encode() for n in offsets[1:])
    return data + f"trailer << /Size 6 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF\n".encode()


def docx_bytes(*, unsupported=False, irregular_table=False):
    from docx import Document
    from docx.oxml import OxmlElement

    document = Document()
    document.add_paragraph("行业分类合成公开样本")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "门类与代码"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "农业"
    table.cell(2, 0).text = "B"
    table.cell(2, 1).text = "采矿业"
    document.add_paragraph("表格之后的解释文字")
    if irregular_table:
        table._tbl.remove(table._tbl.tblGrid)
    if unsupported:
        document.element.body.insert(1, OxmlElement("w:altChunk"))
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
