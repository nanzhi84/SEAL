"""Declared parser boundaries, independent of storage and network services."""

import importlib.util
import unittest
from datetime import datetime
from importlib.metadata import version
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfReader, PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
    NumberObject,
)

from seal.core import SealError, digest
from seal.document_parser import parse_document, parser_version

ROOT = Path(__file__).resolve().parents[2]


def fixture_module(name, path):
    spec = importlib.util.spec_from_file_location(name, ROOT / path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def png_bytes():
    stream = BytesIO()
    Image.new("RGB", (2, 2), "red").save(stream, "PNG")
    return stream.getvalue()


class ParserBoundaries(unittest.TestCase):
    def parse(self, raw, extension, **options):
        snapshot = {
            "url": f"https://example.org/input.{extension}",
            "status": 200,
            "body_hash": digest(raw),
            "headers": {},
            "encoding": None,
        }
        with patch("socket.socket", side_effect=AssertionError("parser accessed network")):
            result = parse_document(snapshot, raw, options=options)
            self.assertEqual(result, parse_document(snapshot, raw, options=options))
        self.assertEqual(set(result), {"format", "children"})
        self.assertLessEqual(sum(n["type"] == "article" for n in result["children"]), 1)
        for node in result["children"]:
            self.assertIn(node["type"], {"article", "table", "image", "json", "document_ref"})
            if node["type"] == "table":
                self.assertLessEqual(len({len(row) for row in node["rows"]}), 1)
        return result

    def rejected(self, raw, extension, code, **options):
        with self.assertRaises(SealError) as raised:
            self.parse(raw, extension, **options)
        self.assertEqual(raised.exception.code, code)

    def test_complete_json_values_and_ambiguous_or_nonfinite_inputs(self):
        examples = (
            (b"null", None),
            (b"false", False),
            (b"22.5", 22.5),
            (b'"source"', "source"),
            (b"[1,null]", [1, None]),
            (b'{"k":{"v":true}}', {"k": {"v": True}}),
        )
        for raw, expected in examples:
            with self.subTest(raw=raw):
                self.assertEqual(
                    self.parse(raw, "json")["children"], [{"type": "json", "value": expected}]
                )
        for raw in (b"NaN", b"Infinity", b"1e999", b'{"k":1,"k":2}', b'{"k":'):
            with self.subTest(raw=raw):
                self.rejected(raw, "json", "json_parse_failed")

    def test_csv_quotes_newlines_empty_positions_and_explicit_delimiter(self):
        result = self.parse(b'\xef\xbb\xbf"a,b",empty\r\n"line1\nline2",\r\n', "csv")
        self.assertEqual(
            result["children"],
            [{"type": "table", "name": "CSV", "rows": [["a,b", "empty"], ["line1\nline2", ""]]}],
        )
        self.assertEqual(
            self.parse(b"A;B\n7;\n", "csv", csv_delimiter=";")["children"][0]["rows"],
            [["A", "B"], ["7", ""]],
        )
        for delimiter in ("\t", " "):
            with self.subTest(delimiter=delimiter):
                raw = f"A{delimiter}B\n7{delimiter}\n".encode()
                self.assertEqual(
                    self.parse(raw, "csv", csv_delimiter=delimiter)["children"][0]["rows"],
                    [["A", "B"], ["7", ""]],
                )
        self.assertEqual(
            self.parse(b"A,B\nC\n", "csv")["children"][0]["rows"],
            [["A", "B"], ["C", None]],
        )
        self.rejected(b'A,"unterminated', "csv", "csv_parse_failed")

    def test_html_listing_selectors_media_order_and_malformed_discovery(self):
        listing = (
            b"<html><head><title>Catalog</title></head><body><main><ul>"
            b'<li><a href="/1">First</a></li><li><a href="/2">Second</a></li>'
            b"</ul></main></body></html>"
        )
        self.assertEqual(
            self.parse(listing, "html", title_css="title")["children"],
            [{"type": "article", "text": "Catalog\nFirst\nSecond", "format": "text"}],
        )
        raw = (
            b'<article><p>Before</p><img src="/figure.png" alt="Plot">'
            b'<table><tr><th colspan="2">Header</th></tr><tr><td>A</td><td>B</td></tr></table>'
            b'<a href="/workbook.xlsx" download>Workbook</a><p>After</p></article>'
        )
        children = self.parse(raw, "html")["children"]
        self.assertEqual(
            [n["type"] for n in children], ["article", "image", "table", "document_ref"]
        )
        self.assertEqual(children[0]["text"], "Before\nWorkbook\nAfter")
        self.assertEqual(
            children[1],
            {"type": "image", "url": "https://example.org/figure.png", "caption": "Plot"},
        )
        self.assertEqual(children[2]["rows"], [["Header", None], ["A", "B"]])
        self.assertEqual(children[3]["url"], "https://example.org/workbook.xlsx")
        self.assertEqual(children[3]["relation"], "attachment")
        malformed = b'<article>Public text<a href="http://[bad" download>Malformed</a></article>'
        self.assertEqual(
            self.parse(malformed, "html")["children"][0]["text"], "Public text Malformed"
        )
        self.rejected(listing, "html", "ambiguous_or_missing_field", body_css="missing")
        self.rejected(listing, "html", "invalid_document_parse_options", body_css="[")
        self.rejected(b'<h1>Login</h1><input type="password">', "html", "access_control_detected")
        self.rejected(b"<h1>404 Not found</h1>", "html", "fallback_error_page")

    def test_docx_embedded_images_merged_table_and_non_body_text(self):
        document = Document()
        document.add_paragraph("Before")
        document.add_picture(BytesIO(png_bytes()))
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged"
        table.cell(1, 0).text = "A"
        table.cell(1, 1).text = "B"
        document.add_paragraph("After")
        document.sections[0].header.paragraphs[0].text = "Header"
        document.sections[0].footer.paragraphs[0].text = "Footer"
        stream = BytesIO()
        document.save(stream)
        raw = stream.getvalue()
        children = self.parse(raw, "docx")["children"]
        self.assertEqual(children[0]["text"], "Before\nAfter\nFooter\nHeader")
        self.assertEqual(children[1]["archive_member"], "word/media/image1.png")
        self.assertEqual(children[1]["body_hash"], digest(raw))
        self.assertEqual(children[2]["rows"], [["Merged", None], ["A", "B"]])
        fixtures = fixture_module(
            "attachment_fixture", "experiments/v1.2-acceptance/attachment_fixtures.py"
        )
        self.rejected(fixtures.docx_bytes(unsupported=True), "docx", "docx_unsupported_content")

    def test_workbooks_keep_sheets_types_merge_gaps_and_source_formulas(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "Main"
        sheet["A1"] = "Title"
        sheet.merge_cells("A1:C2")
        sheet["A3"] = datetime(2026, 10, 11)
        sheet["C3"] = "=2+2"
        book.create_sheet("Empty")
        stream = BytesIO()
        book.save(stream)
        children = self.parse(stream.getvalue(), "xlsx")["children"]
        self.assertEqual([n["name"] for n in children], ["Main", "Empty"])
        self.assertEqual(
            children[0]["rows"],
            [["Title", None, None], [None, None, None], ["2026-10-11T00:00:00", None, "=2+2"]],
        )
        self.assertEqual(children[1]["rows"], [])
        raw = (ROOT / "experiments/v1.2-acceptance/fixtures/public-structured.xls").read_bytes()
        children = self.parse(raw, "xls")["children"]
        self.assertEqual([n["name"] for n in children], ["公开高校", "说明"])
        self.assertEqual(children[0]["rows"][2], ["11001", "示例大学", 42.5, True])
        self.assertEqual(children[0]["rows"][0], ["全国普通高等学校名单", None, None, None])

    def test_pdf_text_and_resource_references_do_not_claim_ocr(self):
        site = fixture_module("pdf_fixture_site", "experiments/v1.1-runtime-acceptance/e2e/site.py")
        reader = PdfReader(BytesIO(site.pdf_bytes()))
        writer = PdfWriter()
        writer.add_page(reader.pages[0])
        image = DecodedStreamObject()
        image.set_data(b"\xff\x00\x00")
        image.update(
            {
                NameObject("/Type"): NameObject("/XObject"),
                NameObject("/Subtype"): NameObject("/Image"),
                NameObject("/Width"): NumberObject(1),
                NameObject("/Height"): NumberObject(1),
                NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
                NameObject("/BitsPerComponent"): NumberObject(8),
            }
        )
        writer.pages[0]["/Resources"][NameObject("/XObject")] = DictionaryObject(
            {NameObject("/Im1"): writer._add_object(image)}
        )
        stream = BytesIO()
        writer.write(stream)
        children = self.parse(stream.getvalue(), "pdf")["children"]
        self.assertEqual(children[0]["text"], "Synthetic PDF notice")
        self.assertEqual(children[1]["archive_member"], "pdf/pages/1/XObject/Im1")
        self.assertEqual(children[1]["body_hash"], digest(stream.getvalue()))
        writer = PdfWriter()
        writer.add_blank_page(width=40, height=40)
        stream = BytesIO()
        writer.write(stream)
        self.rejected(stream.getvalue(), "pdf", "pdf_text_layer_required")
        writer.encrypt("protected")
        stream = BytesIO()
        writer.write(stream)
        self.rejected(stream.getvalue(), "pdf", "pdf_encrypted")

    def test_images_txt_and_failures_have_no_fake_success(self):
        raw = png_bytes()
        self.assertEqual(
            self.parse(raw, "png")["children"],
            [{"type": "image", "url": "https://example.org/input.png", "body_hash": digest(raw)}],
        )
        self.parse(
            b'<svg xmlns="http://www.w3.org/2000/svg"><rect width="1" height="1"/></svg>', "svg"
        )
        self.assertEqual(self.parse(b"A\nB\n", "txt")["children"][0]["text"], "A\nB")
        self.rejected(b"\xff", "txt", "document_decode_failed")
        self.rejected(b" \n", "txt", "empty_document")
        for extension, code in (
            ("pdf", "pdf_parse_failed"),
            ("docx", "docx_parse_failed"),
            ("xls", "xls_parse_failed"),
            ("xlsx", "xlsx_parse_failed"),
            ("png", "image_parse_failed"),
            ("bin", "unsupported_content_type"),
        ):
            with self.subTest(extension=extension):
                self.rejected(b"damaged", extension, code)

    def test_options_input_hash_and_resource_budgets_are_enforced(self):
        for options in (
            {"unexpected": "x"},
            {"format": "auto"},
            {"csv_delimiter": "||"},
            {"body_css": ""},
        ):
            with self.subTest(options=options):
                self.rejected(b"A", "txt", "invalid_document_parse_options", **options)
        snapshot = {"url": "https://example.org/x.txt", "body_hash": digest(b"different")}
        with self.assertRaisesRegex(SealError, "document_body_hash_mismatch"):
            parse_document(snapshot, b"A")
        self.rejected(
            b'<table><tr><td rowspan="999999">x</td></tr></table>',
            "html",
            "document_cell_budget_exceeded",
        )
        with patch("seal.document_parser.MAX_CELLS", 4):
            self.rejected(b"a,b\nc,d\ne,f\n", "csv", "document_cell_budget_exceeded")
        with patch("seal.document_parser.MAX_DOCX_EXPANDED_BYTES", 10):
            self.rejected(b"x" * 11, "txt", "document_input_budget_exceeded")
        stream = BytesIO()
        with ZipFile(stream, "w", compression=ZIP_DEFLATED) as package:
            package.writestr("xl/workbook.xml", " " * 10000)
        with patch("seal.document_parser.MAX_DOCX_EXPANDED_BYTES", 1000):
            self.rejected(stream.getvalue(), "xlsx", "xlsx_expansion_budget_exceeded")

    def test_actual_dependency_changes_change_parser_fingerprint(self):
        current = parser_version()
        self.assertEqual(current, parser_version())
        with patch(
            "seal.document_parser.version", side_effect=lambda name: version(name) + ".changed"
        ):
            self.assertNotEqual(current, parser_version())


if __name__ == "__main__":
    unittest.main()
