"""Public synthetic attachments; no network or production data is used.

The BIFF fixture was generated with xlwt==1.3.0 (a fixture authoring tool,
not a Runtime dependency). Its expected business cells are asserted separately.
"""

from io import BytesIO
from pathlib import Path


def xls_bytes():
    return Path(__file__).with_name("fixtures").joinpath("public-structured.xls").read_bytes()


def docx_bytes(*, unsupported=False, irregular_table=False):
    from docx import Document
    from docx.oxml import OxmlElement

    document = Document()
    document.add_paragraph("行业分类合成公开样本")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "门类与代码"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "农业"
    table.cell(2, 0).text = "B"
    table.cell(2, 1).text = "采矿业"
    document.add_paragraph("表格之后的解释文字")
    if irregular_table:
        table._tbl.remove(table._tbl.tblGrid)
    if unsupported:
        document.element.body.insert(1, OxmlElement("w:altChunk"))
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
