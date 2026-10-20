"""Fixed 16-column workbook encoding and strict local workbook reads."""

import io
import re
import zipfile
from datetime import datetime, timezone

from lxml import etree

from .core import SealError
from .document import COLUMNS

SHEET_TITLE = "公域资料元数据模板"
REQUIRED = ("资料名称", "解析后文件", "资料类型", "文件哈希", "采集时间")
DOCUMENT_TYPES = (
    "法律法规",
    "政策文件",
    "案例",
    "行业标准",
    "通知公告",
    "财务报告",
    "会计准则",
    "税务资料",
    "其他",
)
GUIDANCE = (
    "必填。填写资料正式名称或来源标题。同一资料同一版本包含多个连续分卷时重复填写相同的资料名称"
    "（配合“文件顺序”识别分卷）；同来源下的独立附件作为不同资料填写各自名称。",
    "必填。填写实际上传的 Markdown 文件名并保留 .md（统一按“资料名称_文件哈希前8位.md”重命名，"
    "多卷加序号）；必须与上传文件精确匹配，批次内唯一。一行对应一个解析后文件。",
    "必填。受控值：法律法规、政策文件、案例、行业标准、通知公告、财务报告、会计准则、税务资料、其他。",
    "必填。填写当前行对应单个文件在 OSS 中的完整 SHA-256 哈希值（若无原件则填 Markdown 文件哈希）。"
    "一行对应一个物理文件哈希，用于 API 或 OSS 直接寻址该卷文件。",
    "必填。填写实际采集时间，文本格式 YYYY-MM-DD HH:MM:SS，例如 2026-10-19 10:00:00。",
    "可选。有原件时填写实际交付的原始文件名并保留扩展名（如 .pdf、.docx、.html，统一按"
    "“资料名称_文件哈希前8位.扩展名”重命名防重名）。原件缺失时留空，不编造。",
    "可选。仅用于同一“资料名称”同一版本的连续分卷，填写 1、2…；单文件或独立附件留空。"
    "无法可靠确定顺序时阻止装配，不按 Excel 行序猜测。",
    "可选。按资料类型填写文号、案号、标准号或报告编号。",
    "可选。填写制定机关、发布机构、裁判法院或报告出具机构的正式名称。",
    "可选。填写正式公布、发布或裁判日期，文本格式 YYYY-MM-DD。",
    "可选。仅在资料有明确施行、生效或适用起始日期时填写，文本格式 YYYY-MM-DD；"
    "案例、报告等不适用时留空。",
    "可选。仅在来源明确记载失效或终止日期时填写，文本格式 YYYY-MM-DD；不适用或未核实留空。",
    "可选。填写来源明确记载的状态，如有效、已废止。",
    "可选。填写实际取得资料或元数据的来源名称，如官方数据库、机构网站或资料提供方。",
    "可选。填写可核验的原始发布详情页完整 URL，优先官方来源，不填搜索结果页。",
    "可选。统一填写来源记录的哈希值（如对来源详情页或源记录标识计算的 SHA-256 哈希；"
    "同一个来源记录下可包含多份不同名称的资料或多个分卷）。",
)
WIDTHS = (34, 44, 32, 46, 34, 44, 24, 38, 42, 29, 29, 29, 22, 40, 60, 38)
DATA_START = 3  # Row 1 headers and row 2 field guidance, as in the upload template.
MAX_ROWS = 1048576 - DATA_START + 1
MAX_CELL = 32767
FIXED_TIME = datetime(2000, 1, 1, tzinfo=timezone.utc)
SHA = re.compile(r"[0-9a-f]{64}")
DAY = re.compile(r"\d{4}-\d{2}-\d{2}")
MOMENT = re.compile(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}")


def deterministic_zip(data):
    """Fixed entry timestamps and order so identical inputs give identical bytes."""
    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in sorted(source.infolist(), key=lambda item: item.filename):
            entry = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o600 << 16
            content = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                # openpyxl replaces modified during save, after Workbook properties
                # were frozen. Normalize the serialized field rather than process time.
                root = etree.fromstring(content)
                for node in root:
                    if node.tag in {
                        "{http://purl.org/dc/terms/}created",
                        "{http://purl.org/dc/terms/}modified",
                    }:
                        node.text = "2000-01-01T00:00:00Z"
                content = etree.tostring(root, encoding="utf-8")
            target.writestr(entry, content)
    return output.getvalue()


def workbook_bytes(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if len(rows) > MAX_ROWS:
        raise SealError("metadata_capacity_exceeded")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_TITLE
    header_font, header_fill = Font(bold=True, color="FFFFFFFF"), PatternFill("solid", "FF243B53")
    for index, (name, guidance, width) in enumerate(zip(COLUMNS, GUIDANCE, WIDTHS, strict=True), 1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = width
        for row, value in ((1, name), (2, guidance)):
            cell = sheet.cell(row=row, column=index)
            cell.value = value
            cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.cell(row=1, column=index).font = header_font
        sheet.cell(row=1, column=index).fill = header_fill
    for offset, row in enumerate(rows):
        for index, column in enumerate(COLUMNS, 1):
            value = row[column]
            if len(value) > MAX_CELL:
                raise SealError("metadata_capacity_exceeded")
            cell = sheet.cell(row=DATA_START + offset, column=index)
            try:
                cell.value = value or None
            except ValueError:
                raise SealError("metadata_illegal_character") from None
            if value:
                # Every value is a text cell: official numbers, leading zeros and hashes
                # are preserved and external strings are never evaluated as formulas.
                cell.data_type = "s"
                cell.number_format = "@"
    sheet.freeze_panes = "B3"
    workbook.properties.creator = "SEAL"
    workbook.properties.created = FIXED_TIME.replace(tzinfo=None)
    workbook.properties.modified = FIXED_TIME.replace(tzinfo=None)
    workbook.properties.lastModifiedBy = "SEAL"
    output = io.BytesIO()
    workbook.save(output)
    return deterministic_zip(output.getvalue())


def read_workbook(path):
    """All 16-column rows of a package workbook as strings (empty cells are '')."""
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception:
        raise SealError("metadata_unreadable") from None
    try:
        if workbook.sheetnames != [SHEET_TITLE]:
            raise SealError("metadata_sheet_invalid")
        values = []
        for cells in workbook[SHEET_TITLE].iter_rows():
            if any(cell.data_type == "f" for cell in cells):
                raise SealError("metadata_formula_forbidden")
            values.append(tuple(cell.value for cell in cells))
    finally:
        workbook.close()
    if len(values) < 2 or tuple(values[0][:16]) != COLUMNS or tuple(values[1][:16]) != GUIDANCE:
        raise SealError("metadata_header_invalid")
    rows = []
    for raw in values[DATA_START - 1 :]:
        cells = list(raw) + [None] * (16 - len(raw))
        if any(v is not None for v in cells[16:]):
            raise SealError("metadata_extra_columns")
        if any(v is not None and not isinstance(v, str) for v in cells[:16]):
            raise SealError("metadata_cell_not_text")
        if all(v is None for v in cells[:16]):
            continue
        rows.append({column: cells[i] or "" for i, column in enumerate(COLUMNS)})
    return rows
