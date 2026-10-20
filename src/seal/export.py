"""Read-only file delivery: a fixed Snapshot as metadata.xlsx, Markdown, originals, manifest.

``seal export`` reads only the sealed Snapshot and CAS (never the HTTP API, never a new
acquisition or parse, never the publication pointer). ``seal export verify`` reads only
the local package. A package is staged, fully verified, then atomically moved into place.
"""

import hashlib
import io
import json
import os
import re
import shutil
import signal
import zipfile
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath

from .core import SealError, uid
from .document import COLUMNS, row_values

DELIVERY_SCHEMA = "seal.delivery.v1"
EXPORT_CONTRACT = "seal.export.v1"
SHEET_TITLE = "公域资料元数据模板"
REQUIRED = ("资料名称", "解析后文件", "资料类型", "文件哈希", "采集时间")
DOCUMENT_TYPES = (
    "法律法规",
    "政策文件",
    "案例",
    "行业标准",
    "通知公告",
    "财务报告",
    "会计准则",
    "税务资料",
    "其他",
)
GUIDANCE = (
    "必填。填写资料正式名称或来源标题。同一资料同一版本包含多个连续分卷时重复填写相同的资料名称"
    "（配合“文件顺序”识别分卷）；同来源下的独立附件作为不同资料填写各自名称。",
    "必填。填写实际上传的 Markdown 文件名并保留 .md（统一按“资料名称_文件哈希前8位.md”重命名，"
    "多卷加序号）；必须与上传文件精确匹配，批次内唯一。一行对应一个解析后文件。",
    "必填。受控值：法律法规、政策文件、案例、行业标准、通知公告、财务报告、会计准则、税务资料、其他。",
    "必填。填写当前行对应单个文件在 OSS 中的完整 SHA-256 哈希值（若无原件则填 Markdown 文件哈希）。"
    "一行对应一个物理文件哈希，用于 API 或 OSS 直接寻址该卷文件。",
    "必填。填写实际采集时间，文本格式 YYYY-MM-DD HH:MM:SS，例如 2026-10-19 10:00:00。",
    "可选。有原件时填写实际交付的原始文件名并保留扩展名（如 .pdf、.docx、.html，统一按"
    "“资料名称_文件哈希前8位.扩展名”重命名防重名）。原件缺失时留空，不编造。",
    "可选。仅用于同一“资料名称”同一版本的连续分卷，填写 1、2…；单文件或独立附件留空。"
    "无法可靠确定顺序时阻止装配，不按 Excel 行序猜测。",
    "可选。按资料类型填写文号、案号、标准号或报告编号。",
    "可选。填写制定机关、发布机构、裁判法院或报告出具机构的正式名称。",
    "可选。填写正式公布、发布或裁判日期，文本格式 YYYY-MM-DD。",
    "可选。仅在资料有明确施行、生效或适用起始日期时填写，文本格式 YYYY-MM-DD；"
    "案例、报告等不适用时留空。",
    "可选。仅在来源明确记载失效或终止日期时填写，文本格式 YYYY-MM-DD；不适用或未核实留空。",
    "可选。填写来源明确记载的状态，如有效、已废止。",
    "可选。填写实际取得资料或元数据的来源名称，如官方数据库、机构网站或资料提供方。",
    "可选。填写可核验的原始发布详情页完整 URL，优先官方来源，不填搜索结果页。",
    "可选。统一填写来源记录的哈希值（如对来源详情页或源记录标识计算的 SHA-256 哈希；"
    "同一个来源记录下可包含多份不同名称的资料或多个分卷）。",
)
WIDTHS = (34, 44, 32, 46, 34, 44, 24, 38, 42, 29, 29, 29, 22, 40, 60, 38)
DATA_START = 3  # Row 1 headers and row 2 field guidance, as in the upload template.
MAX_ROWS = 1048576 - DATA_START + 1
MAX_CELL = 32767
FIXED_TIME = datetime(2000, 1, 1, tzinfo=timezone.utc)
SHA = re.compile(r"[0-9a-f]{64}")
DAY = re.compile(r"\d{4}-\d{2}-\d{2}")
MOMENT = re.compile(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}")


class Cancelled(Exception):
    pass


def sha256(data):
    return hashlib.sha256(data).hexdigest()


def file_sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1048576), b""):
            digest.update(chunk)
    return digest.hexdigest()


def safe_relative(value):
    path = PurePosixPath(value)
    if (
        not value
        or path.is_absolute()
        or any(part in ("", ".", "..") for part in path.parts)
        or "\\" in value
        or len(path.parts) > 2
    ):
        raise SealError("delivery_path_invalid")
    return path


def deterministic_zip(data):
    """Fixed entry timestamps and order so identical inputs give identical bytes."""
    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            entry = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o600 << 16
            target.writestr(entry, source.read(info.filename))
    return output.getvalue()


def workbook_bytes(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if len(rows) > MAX_ROWS:
        raise SealError("metadata_capacity_exceeded")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_TITLE
    header_font, header_fill = Font(bold=True, color="FFFFFFFF"), PatternFill("solid", "FF243B53")
    for index, (name, guidance, width) in enumerate(zip(COLUMNS, GUIDANCE, WIDTHS, strict=True), 1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = width
        for row, value in ((1, name), (2, guidance)):
            cell = sheet.cell(row=row, column=index)
            cell.value = value
            cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.cell(row=1, column=index).font = header_font
        sheet.cell(row=1, column=index).fill = header_fill
    for offset, row in enumerate(rows):
        for index, column in enumerate(COLUMNS, 1):
            value = row[column]
            if len(value) > MAX_CELL:
                raise SealError("metadata_capacity_exceeded")
            cell = sheet.cell(row=DATA_START + offset, column=index)
            try:
                cell.value = value or None
            except ValueError:
                raise SealError("metadata_illegal_character") from None
            if value:
                # Every value is a text cell: official numbers, leading zeros and hashes
                # are preserved and external strings are never evaluated as formulas.
                cell.data_type = "s"
                cell.number_format = "@"
    sheet.freeze_panes = "B3"
    workbook.properties.creator = "SEAL"
    workbook.properties.created = FIXED_TIME.replace(tzinfo=None)
    workbook.properties.modified = FIXED_TIME.replace(tzinfo=None)
    workbook.properties.lastModifiedBy = "SEAL"
    output = io.BytesIO()
    workbook.save(output)
    return deterministic_zip(output.getvalue())


def read_workbook(path):
    """All 16-column rows of a package workbook as strings (empty cells are '')."""
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception:
        raise SealError("metadata_unreadable") from None
    try:
        if workbook.sheetnames != [SHEET_TITLE]:
            raise SealError("metadata_sheet_invalid")
        values = list(workbook[SHEET_TITLE].iter_rows(values_only=True))
    finally:
        workbook.close()
    if len(values) < 2 or tuple(values[0][:16]) != COLUMNS or tuple(values[1][:16]) != GUIDANCE:
        raise SealError("metadata_header_invalid")
    rows = []
    for raw in values[DATA_START - 1 :]:
        cells = list(raw) + [None] * (16 - len(raw))
        if any(v is not None for v in cells[16:]):
            raise SealError("metadata_extra_columns")
        if any(v is not None and not isinstance(v, str) for v in cells[:16]):
            raise SealError("metadata_cell_not_text")
        if all(v is None for v in cells[:16]):
            continue
        rows.append({column: cells[i] or "" for i, column in enumerate(COLUMNS)})
    return rows


def check_rows(rows):
    """16-column business rules shared by export staging and offline verify."""
    names, groups = set(), {}
    for row in rows:
        for column in REQUIRED:
            if not row[column]:
                raise SealError("metadata_required_missing")
        if row["资料类型"] not in DOCUMENT_TYPES:
            raise SealError("metadata_document_type_invalid")
        if not SHA.fullmatch(row["文件哈希"]) or (
            row["来源记录号"] and not SHA.fullmatch(row["来源记录号"])
        ):
            raise SealError("metadata_hash_invalid")
        if not MOMENT.fullmatch(row["采集时间"]):
            raise SealError("metadata_time_invalid")
        for column in ("发布日期", "生效日期", "失效日期"):
            if row[column]:
                if not DAY.fullmatch(row[column]):
                    raise SealError("metadata_date_invalid")
                try:
                    datetime.strptime(row[column], "%Y-%m-%d")
                except ValueError:
                    raise SealError("metadata_date_invalid") from None
        if not row["解析后文件"].endswith(".md"):
            raise SealError("metadata_markdown_name_invalid")
        for column in ("解析后文件", "原始文件"):
            if row[column]:
                safe_relative(row[column])
                if "/" in row[column]:
                    raise SealError("delivery_path_invalid")
        if row["解析后文件"] in names:
            raise SealError("metadata_markdown_name_duplicate")
        names.add(row["解析后文件"])
        if row["文件顺序"]:
            if not row["文件顺序"].isdigit() or int(row["文件顺序"]) < 1:
                raise SealError("volume_order_invalid")
            key = (row["资料名称"], row["来源记录号"])
            groups.setdefault(key, []).append(int(row["文件顺序"]))
    for orders in groups.values():
        if sorted(orders) != list(range(1, len(orders) + 1)):
            raise SealError("volume_order_ambiguous")


def check_files(root, rows):
    """Every row has exactly its Markdown and (optional) original with matching SHA-256."""
    expected = set()
    for row in rows:
        markdown = root / "markdown" / row["解析后文件"]
        if not markdown.is_file():
            raise SealError("delivery_file_missing")
        markdown_hash = file_sha256(markdown)
        expected.add(("markdown", row["解析后文件"]))
        if row["原始文件"]:
            original = root / "originals" / row["原始文件"]
            if not original.is_file():
                raise SealError("delivery_file_missing")
            if file_sha256(original) != row["文件哈希"]:
                raise SealError("delivery_file_corrupt")
            expected.add(("originals", row["原始文件"]))
        elif markdown_hash != row["文件哈希"]:
            raise SealError("delivery_file_corrupt")
    actual = set()
    for folder in ("markdown", "originals"):
        path = root / folder
        if path.exists():
            for item in path.iterdir():
                if not item.is_file() or item.is_symlink():
                    raise SealError("delivery_unexpected_file")
                actual.add((folder, item.name))
    if actual != expected:
        raise SealError("delivery_unexpected_file")


def payload_entries(root):
    entries = []
    for path in sorted(root.rglob("*")):
        if path.is_symlink():
            raise SealError("delivery_unexpected_file")
        if path.is_file() and path.name != "manifest.json":
            relative = path.relative_to(root).as_posix()
            entries.append(
                {"path": relative, "bytes": path.stat().st_size, "sha256": file_sha256(path)}
            )
    return entries


def manifest_bytes(value):
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n").encode()


def verify_package(directory):
    """Offline verification: only the local package, never SQL or S3."""
    root = Path(directory)
    try:
        manifest = json.loads((root / "manifest.json").read_bytes())
    except (OSError, ValueError):
        raise SealError("delivery_manifest_unreadable") from None
    if not isinstance(manifest, dict) or manifest.get("schema") != DELIVERY_SCHEMA:
        raise SealError("delivery_manifest_invalid")
    files = manifest.get("files")
    if not isinstance(files, list):
        raise SealError("delivery_manifest_invalid")
    for item in files:
        safe_relative(item.get("path", ""))
    entries = payload_entries(root)
    if entries != sorted(files, key=lambda item: item["path"]):
        raise SealError("delivery_payload_mismatch")
    rows = read_workbook(root / "metadata.xlsx")
    check_rows(rows)
    check_files(root, rows)
    lineage = manifest.get("rows")
    counts = manifest.get("counts") or {}
    if (
        not isinstance(lineage, list)
        or len(lineage) != len(rows)
        or counts.get("rows") != len(rows)
        or counts.get("markdown_files") != len({r["解析后文件"] for r in rows})
        or counts.get("original_files") != len({r["原始文件"] for r in rows if r["原始文件"]})
    ):
        raise SealError("delivery_manifest_mismatch")
    for number, (row, item) in enumerate(zip(rows, lineage, strict=True), DATA_START):
        if (
            item.get("row") != number
            or item.get("markdown_file") != "markdown/" + row["解析后文件"]
            or item.get("file_hash") != row["文件哈希"]
            or item.get("source_record_hash") != row["来源记录号"]
            or (item.get("original_file") or "")
            != ("originals/" + row["原始文件"] if row["原始文件"] else "")
            or (str(item.get("volume_order") or "")) != row["文件顺序"]
        ):
            raise SealError("delivery_manifest_mismatch")
        if file_sha256(root / "markdown" / row["解析后文件"]) != item.get("markdown_hash"):
            raise SealError("delivery_manifest_mismatch")
    return {
        "status": "verified",
        "schema": DELIVERY_SCHEMA,
        "snapshot_id": manifest.get("snapshot_id"),
        "sources": manifest.get("sources"),
        "rows": len(rows),
        "files": len(entries),
    }


def collect_rows(snapshot_id, sources):
    from .db import connect
    from .public_results import snapshot_rows
    from .snapshots import header

    with connect() as c:
        c.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ, READ ONLY")
        snapshot = header(c, snapshot_id)
        frozen = {
            r["source_id"]: r["coverage"]
            for r in c.execute(
                "SELECT source_id,coverage FROM seal.snapshot_source WHERE snapshot_id=%s",
                (snapshot_id,),
            )
        }
        selected = sorted(set(sources)) if sources else sorted(frozen)
        if not selected or any(s not in frozen for s in selected):
            raise SealError("export_source_not_in_snapshot")
        rows, after = [], ("", "", 0)
        while True:
            page = snapshot_rows(c, snapshot_id, selected, after, 1000)
            rows.extend(page)
            if len(page) < 1000:
                break
            last = page[-1]
            after = (last["source_id"], last["resource_id"], last["row_number"])
    return snapshot, selected, {s: frozen[s] for s in selected}, rows


def export(snapshot_id, sources, output):
    """Materialize a sealed Snapshot into a verified, atomically completed package."""
    from .objects import Objects, work_directory
    from .snapshots import identity

    identity(snapshot_id)
    output = Path(output).resolve()
    if output.exists() and not output.is_dir():
        raise SealError("export_destination_conflict")
    snapshot, selected, coverage, rows = collect_rows(snapshot_id, sources)
    if len(rows) > MAX_ROWS:
        raise SealError("metadata_capacity_exceeded")
    staging = work_directory() / "exports" / uid()
    cancelled = []
    previous = {}

    def cancel(signum, frame):
        cancelled.append(signum)

    for sig in (signal.SIGTERM, signal.SIGINT):
        previous[sig] = signal.signal(sig, cancel)

    def active():
        if cancelled:
            raise SealError("interrupted")

    try:
        (staging / "markdown").mkdir(parents=True)
        (staging / "originals").mkdir()
        objects = Objects()
        written = {}
        table, lineage, resources = [], [], set()

        def put(folder, name, ref):
            key = (folder, name)
            if key in written:
                if written[key] != ref:
                    raise SealError("delivery_filename_collision")
                return
            data = objects.get(ref)
            path = staging / folder / name
            with open(path, "xb") as handle:
                handle.write(data)
            written[key] = ref

        for number, row in enumerate(rows, DATA_START):
            active()
            values = row_values(row)
            put("markdown", row["markdown_file"], row["markdown_hash"])
            if row["original_ref"]:
                put("originals", row["original_file"], row["original_ref"])
            table.append(values)
            resources.add((row["source_id"], row["resource_id"]))
            lineage.append(
                {
                    "row": number,
                    "source_id": row["source_id"],
                    "resource_id": row["resource_id"],
                    "version_id": row["version_id"],
                    "source_record_hash": row["source_record_hash"],
                    "markdown_file": "markdown/" + row["markdown_file"],
                    "markdown_hash": row["markdown_hash"],
                    "original_file": "originals/" + row["original_file"]
                    if row["original_file"]
                    else None,
                    "file_hash": row["file_hash"],
                    "role": row["role"],
                    "volume_order": row["volume_order"],
                    "locator": row["locator"],
                }
            )
        active()
        (staging / "metadata.xlsx").write_bytes(workbook_bytes(table))
        if not any((staging / "originals").iterdir()):
            (staging / "originals").rmdir()
        # Reopen the workbook and recheck every row, file and hash before completion.
        reread = read_workbook(staging / "metadata.xlsx")
        if reread != table:
            raise SealError("metadata_roundtrip_mismatch")
        check_rows(reread)
        check_files(staging, reread)
        active()
        manifest = {
            "schema": DELIVERY_SCHEMA,
            "export_contract": EXPORT_CONTRACT,
            "snapshot_id": snapshot_id,
            "snapshot_created_at": snapshot["created_at"].isoformat(),
            "sources": selected,
            "coverage": coverage,
            "counts": {
                "documents": len(resources),
                "rows": len(table),
                "markdown_files": len({r["解析后文件"] for r in table}),
                "original_files": len({r["原始文件"] for r in table if r["原始文件"]}),
            },
            "columns": list(COLUMNS),
            "data_start_row": DATA_START,
            "rows": lineage,
            "files": payload_entries(staging),
        }
        (staging / "manifest.json").write_bytes(manifest_bytes(manifest))
        verify_package(staging)
        active()
        if output.exists():
            if not (output / "manifest.json").is_file() or (
                (output / "manifest.json").read_bytes() != (staging / "manifest.json").read_bytes()
            ):
                raise SealError("export_destination_conflict")
            verify_package(output)
            return dict(verify_package(staging), status="unchanged", output=str(output))
        output.parent.mkdir(parents=True, exist_ok=True)
        try:
            os.rename(staging, output)
        except OSError as exc:
            if exc.errno != 18:  # EXDEV: stage beside the destination, then rename.
                raise SealError("export_move_failed") from None
            sibling = output.parent / (".staging-" + uid())
            shutil.copytree(staging, sibling)
            try:
                os.rename(sibling, output)
            except OSError:
                shutil.rmtree(sibling, ignore_errors=True)
                raise SealError("export_move_failed") from None
        return dict(verify_package(output), status="exported", output=str(output))
    finally:
        for sig, handler in previous.items():
            signal.signal(sig, handler)
        shutil.rmtree(staging, ignore_errors=True)
