"""Bounded projections of archived responses into the five-node document contract."""

import csv
import json
import math
import warnings
from io import BytesIO, StringIO
from urllib.parse import urljoin, urlsplit
from zipfile import ZipFile

from lxml import html
from PIL import Image
from pypdf import PdfReader

from .attachments import MAX_CELLS, docx_structure, docx_text, xls_structure
from .core import SealError, public_url
from .fallback import _check_page
from .scope import paths_for

PARSER_VERSION = "document-ir.v1"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
IMAGE_MIMES = {"image/png", "image/jpeg", "image/gif", "image/webp"}


def _decode(body):
    try:
        return body.decode("utf-8-sig")
    except UnicodeError:
        raise SealError("document_text_encoding_unsupported") from None


def _scalar(value):
    from datetime import date, datetime, time
    from decimal import Decimal

    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if value is None or type(value) in (str, int, float, bool):
        if type(value) is float and not math.isfinite(value):
            raise SealError("document_cell_unsupported")
        return value
    raise SealError("document_cell_unsupported")


def _table(name, rows):
    if sum(len(row) for row in rows) > MAX_CELLS:
        raise SealError("document_table_too_large")
    width = max((len(row) for row in rows), default=0)
    return {
        "type": "table",
        "name": name,
        "columns": None,
        "rows": [[_scalar(cell) for cell in row] + [None] * (width - len(row)) for row in rows],
    }


def _xlsx(body):
    from openpyxl import load_workbook

    try:
        with ZipFile(BytesIO(body)) as package:
            entries = package.infolist()
            if len(entries) > 2000 or sum(entry.file_size for entry in entries) > 64 * 1024 * 1024:
                raise SealError("xlsx_expansion_budget_exceeded")
        book = load_workbook(BytesIO(body), read_only=True, data_only=True)
        result = []
        total = 0
        for sheet in book:
            if sheet.max_row and sheet.max_column and sheet.max_row * sheet.max_column > MAX_CELLS:
                raise SealError("document_table_too_large")
            rows = []
            for row in sheet.iter_rows(values_only=True):
                total += len(row)
                if total > MAX_CELLS:
                    raise SealError("document_table_too_large")
                rows.append(list(row))
            result.append(_table(sheet.title, rows))
        book.close()
        return result
    except SealError:
        raise
    except Exception:
        raise SealError("xlsx_parse_failed") from None


def _docx(body):
    blocks = docx_structure(body)
    children, paragraphs, article_index = [], [], None
    for block in blocks:
        if "text" in block:
            if block["text"].strip():
                if article_index is None:
                    article_index = len(children)
                    children.append(None)
                paragraphs.append(block["text"].strip())
        else:
            rows = [[docx_text(cell["blocks"]) for cell in row["cells"]] for row in block["rows"]]
            children.append(_table(None, rows))
    if article_index is not None:
        children[article_index] = {"type": "article", "text": "\n".join(paragraphs)}
    return children


def _html(body, base_url, config, selector=None):
    try:
        decoded = _decode(body)
        root = html.fromstring(decoded)
    except Exception:
        raise SealError("html_parse_failed") from None
    for node in root.xpath("//script|//style|//nav|//footer|//form|//noscript"):
        node.drop_tree()
    if selector is not None:
        if not isinstance(selector, str) or not 0 < len(selector) <= 512:
            raise SealError("document_content_selector_invalid")
        try:
            selected = root.cssselect(selector)
        except Exception:
            raise SealError("document_content_selector_invalid") from None
        if len(selected) != 1:
            raise SealError("document_content_selector_invalid")
        scope = selected[0]
    else:
        scope = next(iter(root.xpath("//article|//main|//*[@role='main']|//body")), root)
        if scope.tag == "body" and not scope.xpath(".//a[@href]"):
            # A list page may keep its links in <nav>, with or without a
            # heading outside it. The Recipe can follow those links too.
            original = html.fromstring(decoded)
            for node in original.xpath("//script|//style|//footer|//form|//noscript"):
                node.drop_tree()
            original_body = next(iter(original.xpath("//body")), original)
            if original_body.xpath(".//nav//a[@href]"):
                scope = original_body
    children = []
    text = " ".join(" ".join(scope.itertext()).split())
    if text:
        children.append({"type": "article", "text": text})
    seen = set()
    for node in scope.iterdescendants():
        if node.tag == "table":
            rows = [
                [" ".join(cell.itertext()).strip() for cell in tr.xpath("./th|./td")]
                for tr in node.xpath(".//tr")
            ]
            if rows:
                caption = " ".join(node.xpath("./caption//text()")).strip() or None
                children.append(_table(caption, rows))
        elif node.tag == "img" and node.get("src"):
            try:
                target = public_url(urljoin(base_url, node.get("src")))
            except SealError:
                continue
            children.append({"type": "image", "url": target, "caption": node.get("alt") or None})
        elif node.tag == "a" and node.get("href"):
            try:
                target = public_url(urljoin(base_url, node.get("href")))
            except SealError:
                continue
            parsed = urlsplit(target)
            if not paths_for(config, parsed.hostname, parsed.path) or target in seen:
                continue
            seen.add(target)
            children.append(
                {
                    "type": "document_ref",
                    "url": target,
                    "name": " ".join(node.itertext()).strip() or None,
                }
            )
    return children


def _kind(response):
    mime = response.headers.get("Content-Type", b"").decode("latin1").split(";", 1)[0].lower()
    suffix = urlsplit(response.url).path.lower()
    if (
        mime.startswith("text/html")
        or suffix.endswith((".html", ".htm"))
        or response.body.lstrip().lower().startswith((b"<html", b"<!doctype html"))
    ):
        return "html", mime
    if mime in IMAGE_MIMES or suffix.endswith((".png", ".jpg", ".jpeg", ".gif", ".webp")):
        return "image", mime
    if mime == "application/pdf" or suffix.endswith(".pdf"):
        return "pdf", mime
    if mime == XLSX_MIME or suffix.endswith(".xlsx"):
        return "xlsx", mime
    if mime == DOCX_MIME or suffix.endswith(".docx"):
        return "docx", mime
    if mime == "application/vnd.ms-excel" or suffix.endswith(".xls"):
        return "xls", mime
    if mime in {"text/csv", "application/csv"} or suffix.endswith(".csv"):
        return "csv", mime
    if mime in {"application/json", "text/json"} or suffix.endswith(".json"):
        return "json", mime
    if mime == "text/plain" or suffix.endswith(".txt"):
        return "txt", mime
    raise SealError("document_format_unsupported")


def parse_document(response, config, snapshot_id, body_hash):
    if response.status != 200 or response.request.method not in {"GET", "POST"}:
        raise SealError("document_response_ineligible")
    kind, mime = _kind(response)
    body = response.body
    if kind == "html":
        _check_page(response)
        children = _html(
            body,
            response.urljoin(""),
            config,
            response.meta.get("seal_document_content_css"),
        )
    elif kind == "pdf":
        try:
            reader = PdfReader(BytesIO(body), strict=True)
            text = "\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()
        except Exception:
            raise SealError("pdf_parse_failed") from None
        if not text:
            raise SealError("pdf_text_layer_required")
        children = [{"type": "article", "text": text}]
    elif kind == "docx":
        children = _docx(body)
    elif kind == "xls":
        children = [_table(sheet["name"], sheet["rows"]) for sheet in xls_structure(body)]
    elif kind == "xlsx":
        children = _xlsx(body)
    elif kind == "csv":
        try:
            rows = list(csv.reader(StringIO(_decode(body), newline=""), strict=True))
        except csv.Error:
            raise SealError("csv_parse_failed") from None
        children = [_table(None, rows)]
    elif kind == "json":
        try:
            value = json.loads(
                _decode(body), parse_constant=lambda _: (_ for _ in ()).throw(ValueError())
            )
        except (ValueError, RecursionError):
            raise SealError("json_parse_failed") from None
        children = [{"type": "json", "value": value}]
    elif kind == "txt":
        children = [{"type": "article", "text": _decode(body)}]
    else:
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(BytesIO(body)) as image:
                    image_format = image.format
                    if image.width * image.height > 100_000_000:
                        raise ValueError
                    image.verify()
        except Exception:
            raise SealError("image_parse_failed") from None
        actual_mime = {
            "PNG": "image/png",
            "JPEG": "image/jpeg",
            "GIF": "image/gif",
            "WEBP": "image/webp",
        }.get(image_format)
        suffix = urlsplit(response.url).path.lower().rsplit(".", 1)[-1]
        suffix_mime = {
            "png": "image/png",
            "jpg": "image/jpeg",
            "jpeg": "image/jpeg",
            "gif": "image/gif",
            "webp": "image/webp",
        }.get(suffix)
        if (
            actual_mime is None
            or mime in IMAGE_MIMES
            and mime != actual_mime
            or (suffix_mime is not None and suffix_mime != actual_mime)
        ):
            raise SealError("image_format_mismatch")
        mime = actual_mime
        children = [
            {
                "type": "image",
                "url": public_url(response.url),
                "snapshot_id": snapshot_id,
                "body_hash": body_hash,
                "caption": None,
            }
        ]
    if not children and kind != "html":
        raise SealError("empty_document")
    return {"type": "document", "children": children}, kind, mime
