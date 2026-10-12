"""Office format projections; the common parser owns dispatch and validation."""

import math
from datetime import date, datetime, time, timedelta
from io import BytesIO

import openpyxl
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from lxml import etree

from .attachments import MAX_CELLS, _docx_document
from .core import SealError


def _docx_table(table):
    rows, width = [], len(table.columns)
    if len(table.rows) * width > MAX_CELLS:
        raise SealError("document_cell_budget_exceeded")
    for row in table.rows:
        values = [None] * row.grid_cols_before
        for node in row._tr.tc_lst:
            cell = _Cell(node, table)
            text = "\n".join(
                part.text
                if isinstance(part, Paragraph)
                else "\n".join("\t".join(c.text for c in r.cells) for r in part.rows)
                for part in cell.iter_inner_content()
            )
            values.append(None if node.vMerge == "continue" else text)
            values.extend([None] * (node.grid_span - 1))
        values.extend([None] * row.grid_cols_after)
        rows.append(values)
    return rows


def _docx(snapshot, body):
    from .document_parser import _Content, _image_reference, _target, _zip_budget

    try:
        document = _docx_document(body)
        content, table_count = _Content(), 0
        unsupported = {qn("w:" + name) for name in ("altChunk", "object", "subDoc")}
        if any(node.tag in unsupported for node in document.element.body.iter()):
            raise SealError("docx_unsupported_content")

        def images(element, part):
            for blip in element.iter(qn("a:blip")):
                rel_id = blip.get(qn("r:embed")) or blip.get(qn("r:link"))
                if not rel_id or rel_id not in part.rels:
                    raise SealError("docx_image_relationship_missing")
                relationship = part.rels[rel_id]
                if relationship.is_external:
                    url = _target(snapshot["url"], relationship.target_ref)
                    if url:
                        content.children.append({"type": "image", "url": url})
                else:
                    member = str(relationship.target_part.partname).lstrip("/")
                    content.children.append(_image_reference(snapshot, member))

        for part in document.iter_inner_content():
            if isinstance(part, Paragraph):
                content.text(part.text)
                images(part._p, document.part)
            elif isinstance(part, Table):
                table_count += 1
                content.table(f"Table {table_count}", _docx_table(part))
                images(part._tbl, document.part)
        # Keep text outside the main body without pretending to restore layout.
        with _zip_budget(body, "docx") as package:
            by_name = {str(part.partname).lstrip("/"): part for part in document.part.package.parts}
            for path in sorted(package.namelist()):
                if not (
                    path.startswith("word/")
                    and path.endswith(".xml")
                    and path.split("/")[-1].startswith(
                        ("header", "footer", "footnotes", "endnotes")
                    )
                ):
                    continue
                root = etree.fromstring(
                    package.read(path), etree.XMLParser(resolve_entities=False, no_network=True)
                )
                for paragraph in root.iter(qn("w:p")):
                    content.text("".join(node.text or "" for node in paragraph.iter(qn("w:t"))))
                if path in by_name:
                    images(root, by_name[path])
        return content.finish()
    except SealError:
        raise
    except Exception:
        raise SealError("docx_parse_failed") from None


def _scalar(value):
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise SealError("xlsx_nonfinite_number")
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    raise SealError("xlsx_cell_type_unsupported")


def _xlsx(body):
    book = None
    from .document_parser import _Content, _zip_budget

    try:
        with _zip_budget(body, "xlsx"):
            pass
        book = openpyxl.load_workbook(
            BytesIO(body), data_only=False, read_only=False, keep_links=False
        )
        content = _Content()
        for sheet in book.worksheets:
            if sheet.max_row * sheet.max_column > MAX_CELLS:
                raise SealError("document_cell_budget_exceeded")
            used = [
                (cell.row, cell.column)
                for row in sheet.iter_rows()
                for cell in row
                if cell.value is not None
            ]
            used.extend((merged.max_row, merged.max_col) for merged in sheet.merged_cells.ranges)
            height = max((r for r, _ in used), default=0)
            width = max((c for _, c in used), default=0)
            rows = []
            for row in (
                sheet.iter_rows(min_row=1, max_row=height, min_col=1, max_col=width) if used else ()
            ):
                if any(cell.data_type == "e" for cell in row):
                    raise SealError("xlsx_cell_error")
                rows.append([_scalar(cell.value) for cell in row])
            content.table(sheet.title, rows)
        return content.finish()
    except SealError:
        raise
    except Exception:
        raise SealError("xlsx_parse_failed") from None
    finally:
        if book is not None:
            book.close()
