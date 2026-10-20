"""Independent issue #32 CLI acceptance infrastructure; never imports seal internals."""
from __future__ import annotations
import hashlib, io, json, os, secrets, socket, subprocess, sys, threading, time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.request import Request, urlopen
import boto3, psycopg, yaml
from psycopg import sql
from psycopg.conninfo import conninfo_to_dict, make_conninfo
from docx import Document
from cryptography.fernet import Fernet

ROOT = Path(__file__).resolve().parents[1]
PYTHON = Path(sys.executable)
HEADERS = ['资料名称','解析后文件','资料类型','文件哈希','采集时间','原始文件','文件顺序','文号/案号','发布机构','发布日期','生效日期','失效日期','状态','来源名称','来源详情页','来源记录号']

def sha(data): return hashlib.sha256(data).hexdigest()
def port():
    with socket.socket() as sock:
        sock.bind(('127.0.0.1',0)); return sock.getsockname()[1]
def docx(text):
    doc=Document(); doc.add_heading(text,0)
    doc.add_paragraph(('This official acceptance fixture contains substantive provisions and auditable evidence. '+text+' ')*30)
    out=io.BytesIO(); doc.save(out); return out.getvalue()

class Lab:
    def __init__(self):
        self.name='issue32-'+time.strftime('%Y%m%d%H%M%S')+'-'+secrets.token_hex(3)
        self.work=ROOT/'work'/self.name; self.out=ROOT/'artifacts'/self.name
        self.work.mkdir(parents=True); self.out.mkdir(parents=True)
        self.env=os.environ.copy(); self.env['PYTHONPATH']=str(ROOT/'src')
        self.env['SEAL_ALLOW_LOOPBACK']='1'; self.env['SEAL_WORK_DIR']=str(self.work/'runtime')
        old=conninfo_to_dict(self.env['SEAL_DATABASE_URL']); self.db='accept32_'+secrets.token_hex(7)
        with psycopg.connect(self.env['SEAL_DATABASE_URL'],autocommit=True) as con:
            con.execute(sql.SQL('CREATE DATABASE {}').format(sql.Identifier(self.db)))
        old['dbname']=self.db; self.env['SEAL_DATABASE_URL']=make_conninfo(**old)
        self.bucket='accept32-'+secrets.token_hex(10); self.env['SEAL_S3_BUCKET']=self.bucket
        self.s3=boto3.client('s3',endpoint_url=self.env['SEAL_S3_ENDPOINT'],aws_access_key_id=self.env['SEAL_S3_ACCESS_KEY'],aws_secret_access_key=self.env['SEAL_S3_SECRET_KEY'],region_name=self.env.get('SEAL_S3_REGION','us-east-1'))
        self.s3.create_bucket(Bucket=self.bucket)
        self.log=[]; self.checks=[]; self.children=[]; self.files={}; self.requests=[]
        self.server_port=port(); self.base=f'http://127.0.0.1:{self.server_port}'
        lab=self
        class Handler(BaseHTTPRequestHandler):
            def do_GET(self):
                lab.requests.append(self.path)
                item=lab.files.get(self.path)
                if item is None: self.send_response(404); self.end_headers(); return
                media,body=item; self.send_response(200); self.send_header('Content-Type',media)
                self.send_header('Content-Length',str(len(body))); self.end_headers(); self.wfile.write(body)
            def log_message(self,*args): pass
        self.http=ThreadingHTTPServer(('127.0.0.1',self.server_port),Handler)
        threading.Thread(target=self.http.serve_forever,daemon=True).start()
        self.populate()
    def populate(self):
        self.files['/robots.txt']=('text/plain',b'User-agent: *\nAllow: /\n')
        links=''.join(f'<a href="/detail/{x}">{x}</a>' for x in ('single','volumes','attachments'))
        self.files['/index']=('text/html; charset=utf-8',f'<html><body>{links}</body></html>'.encode())
        pages={
          'single':('<h1>单件法规</h1><p>'+'Concrete single document provisions and reliable acceptance text. '*60+'</p>'),
          'volumes':'<h1>连续分卷法规</h1><a href="/files/v1.docx">Part 1</a><a href="/files/v2.docx">Part 2</a>',
          'attachments':'<h1>附件主资料</h1><a href="/files/main.docx">主文件</a><a href="/files/note.docx">独立修订说明</a>',
          'ambiguous':'<h1>模糊分卷</h1><a href="/files/v1.docx">Part 1</a><a href="/files/v2.docx">Part 1</a>',
          'later':'<h1>后续快照资料</h1><p>'+'Later snapshot B content has an independent explicit resource. '*60+'</p>',
        }
        for key,body in pages.items(): self.files['/detail/'+key]=('text/html; charset=utf-8',f'<html><body>{body}</body></html>'.encode())
        for key in ('v1','v2','main','note'): self.files['/files/'+key+'.docx']=('application/vnd.openxmlformats-officedocument.wordprocessingml.document',docx('Acceptance '+key))
    def config(self,entry='/index',name='source',title=None,source='dd-017'):
        fields={'document_number':{'value':'000042'},'issuer':{'value':'=SUM(1,2)'},'published_at':{'value':'2026-10-01'}}
        if title is not None: fields['title']={'value':title}
        body={'id':source,'adapter':'web','entry_urls':[self.base+entry],'allowed_hosts':['127.0.0.1'],'allowed_path_prefixes':['/'],'archive_approved':True,'scope':'isolated issue32 business acceptance','robots':True,'concurrency':2,'delay':0,'recheck_seconds':0,'budget':{'requests':60,'seconds':90,'response_bytes':1048576,'documents':100},'fetch':{'mode':'http','max_depth':2,'retries':0},'metadata':{'source_name':'+official fixture','document_type':'法律法规','language':'en'},'documents':{'detail_url_patterns':['/detail/'],'detail_body':'auto','fields':fields}}
        path=self.work/(name+'.yaml'); path.write_text(yaml.safe_dump(body,allow_unicode=True)); return path
    def cli(self,*args,ok=True,env=None,timeout=100):
        command=[str(PYTHON),'-m','seal',*map(str,args)]
        result=subprocess.run(command,cwd=ROOT,env=env or self.env,text=True,capture_output=True,timeout=timeout)
        entry={'args':list(map(str,args)),'exit_code':result.returncode,'stdout':result.stdout,'stderr':result.stderr}
        self.log.append(entry); (self.out/'commands.json').write_text(json.dumps(self.log,ensure_ascii=False,indent=2))
        if ok and result.returncode: raise AssertionError(f'CLI {args[0]} exited {result.returncode}: {result.stderr[-1800:]} {result.stdout[-1800:]}')
        if ok is False and not result.returncode: raise AssertionError(f'CLI unexpectedly succeeded: {args}')
        return result
    def state(self):
        tables={}
        with psycopg.connect(self.env['SEAL_DATABASE_URL']) as con:
            con.execute('SET TRANSACTION READ ONLY')
            names=[r[0] for r in con.execute("select tablename from pg_tables where schemaname='seal' order by tablename")]
            for name in names:
                rows=[r[0] for r in con.execute(sql.SQL('select row_to_json(t)::text from seal.{} t order by row_to_json(t)::text').format(sql.Identifier(name)))]
                tables[name]={'count':len(rows),'sha256':sha('\n'.join(rows).encode())}
        objects=[]
        for page in self.s3.get_paginator('list_objects_v2').paginate(Bucket=self.bucket):
            for obj in page.get('Contents',[]): objects.append((obj['Key'],obj['Size'],obj['ETag']))
        return {'tables':tables,'objects':sorted(objects)}
    def check(self,name,condition,evidence=None,fatal=True):
        self.checks.append({'name':name,'passed':bool(condition),'evidence':evidence})
        if not condition and fatal: raise AssertionError(name)
    def snapshot(self):
        result=self.cli('snapshot','create'); data=json.loads(result.stdout)
        if isinstance(data,str): return data
        return data.get('snapshot_id') or data.get('id') or data['snapshot']['id']
    def api(self):
        self.token=secrets.token_urlsafe(32); self.api_port=port(); self.api_base=f'http://127.0.0.1:{self.api_port}'
        body={'public_base_url':self.api_base,'token_key':Fernet.generate_key().decode(),'request_timeout_seconds':60,'file_max_bytes':134217728,'callers':[{'id':'accept32','key_sha256':sha(self.token.encode()),'sources':['dd-017']}]}
        conf=self.work/'api.private.json'; conf.write_text(json.dumps(body)); conf.chmod(0o600)
        log=(self.work/'api.log').open('w'); child=subprocess.Popen([str(PYTHON),'-m','seal','serve','--config',str(conf),'--host','127.0.0.1','--port',str(self.api_port)],cwd=ROOT,env=self.env,stdout=log,stderr=log)
        self.children.append(child)
        for _ in range(100):
            try:
                with urlopen(self.api_base+'/healthz',timeout=1) as response:
                    if response.status==200: return
            except OSError: time.sleep(.1)
        raise AssertionError('API failed to start; inspect private work log')
    def get(self,path):
        with urlopen(Request(self.api_base+path,headers={'Authorization':'Bearer '+self.token}),timeout=30) as response: return response.read()
    def finish(self,error=None):
        for proc in self.children:
            proc.terminate()
            try: proc.wait(timeout=5)
            except subprocess.TimeoutExpired: proc.kill()
        self.http.shutdown()
        hashes={str(p.relative_to(ROOT)):sha(p.read_bytes()) for folder in ('src','tests') for p in sorted((ROOT/folder).rglob('*')) if p.is_file() and '__pycache__' not in p.parts}
        (self.out/'source-test-hashes.json').write_text(json.dumps(hashes,indent=2,ensure_ascii=False))
        report={'run':self.name,'head':subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),'checks':self.checks,'error':error,'fixture_requests':self.requests,'tests_sha256':{k:v for k,v in hashes.items() if k.startswith('tests/issue32_')},'environment':{'python':sys.version.split()[0],'database':'isolated acceptance-owned PostgreSQL database','objects':'isolated acceptance-owned S3 bucket','inputs':'loopback HTTP HTML and generated DOCX','secret_values_recorded':False}}
        (self.out/'report.json').write_text(json.dumps(report,indent=2,ensure_ascii=False))
        print(json.dumps({'artifacts':str(self.out),'checks':len(self.checks),'passed':sum(c['passed'] for c in self.checks),'error':error},ensure_ascii=False))
