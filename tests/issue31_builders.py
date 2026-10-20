"""Issue #31 independent acceptance support: isolated services, fixtures, evidence.

Nothing here imports ``seal``; the product is exercised only through its real CLI
(subprocess), its PostgreSQL ``seal`` schema, the S3 API and its HTTP API.
"""

from __future__ import annotations

import contextlib
import hashlib
import io
import json
import os
import re
import secrets
import shutil
import signal
import socket
import subprocess
import sys
import threading
import time
import traceback
import uuid
import zipfile
from datetime import datetime, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlsplit

import boto3
import psycopg
import requests
import yaml
from psycopg.rows import dict_row

from issue31_primitives import NotVerified

# --------------------------------------------------------------------------- document builders

TCVN3_SOURCE = "µ¸¶·¹¨»¾¼½Æ©ÇÊÈÉË®ÌÐÎÏÑªÒÕÓÔÖ×ÝØÜÞßãáâä«åèæçé¬êíëìîïóñòô-õøö÷ùúýûüþ¡¢§£¤¥¦"
TCVN3_TARGET = "àáảãạăằắẳẵặâầấẩẫậđèéẻẽẹêềếểễệìíỉĩịòóỏõọôồốổỗộơờớởỡợùúủũụưừứửữựỳýỷỹỵĂÂĐÊÔƠƯ"
_ENC = {t: ("\u00ad" if s == "-" else s) for s, t in zip(TCVN3_SOURCE, TCVN3_TARGET)}  # ư is soft hyphen 0xAD


def to_tcvn3(text: str) -> str:
    out = []
    for ch in text:
        if ch in _ENC:
            out.append(_ENC[ch])
        elif ord(ch) < 128:
            out.append(ch)
        else:
            raise ValueError(f"not TCVN3-encodable: {ch!r}")
    return "".join(out)


def rtf_escape_latin1(text: str) -> str:
    out = []
    for ch in text:
        code = ord(ch)
        if ch in "\\{}":
            out.append("\\" + ch)
        elif code < 128:
            out.append(ch)
        elif code < 256:
            out.append(f"\\'{code:02x}")
        else:
            raise ValueError(ch)
    return "".join(out)


def rtf_escape_unicode(text: str) -> str:
    out = []
    for ch in text:
        code = ord(ch)
        if ch in "\\{}":
            out.append("\\" + ch)
        elif code < 128:
            out.append(ch)
        else:
            signed = code if code < 32768 else code - 65536
            out.append(f"\\u{signed}?")
    return "".join(out)


def tcvn3_rtf(unicode_lines: list[str], tcvn3_lines: list[str], unicode_tail: list[str] = ()) -> bytes:
    """RTF with Unicode header lines (Times New Roman, \\uN) and TCVN3 body (.VnTime)."""
    parts = [
        r"{\rtf1\ansi\ansicpg1252\deff0",
        r"{\fonttbl{\f0\froman\fcharset0 Times New Roman;}{\f1\fnil\fcharset0 .VnTime;}{\f2\fnil\fcharset0 .VnTimeH;}}",
        r"{\stylesheet{\s0 Normal;}}",
        r"{\header \pard\plain HEADERNOISE-SHOULD-NOT-APPEAR\par}",
        r"{\footer \pard\plain FOOTERNOISE-SHOULD-NOT-APPEAR\par}",
    ]
    for line in unicode_lines:
        parts.append(r"\pard\plain\f0\fs24 " + rtf_escape_unicode(line) + r"\par")
    for line in tcvn3_lines:
        parts.append(r"\pard\plain\f1\fs24 " + rtf_escape_latin1(to_tcvn3(line)) + r"\par")
    for line in unicode_tail:
        parts.append(r"\pard\plain\f0\fs24 " + rtf_escape_unicode(line) + r"\par")
    parts.append("}")
    return "\n".join(parts).encode("ascii")


def plain_rtf(lines: list[str]) -> bytes:
    body = "\n".join(r"\pard\plain\f0 " + rtf_escape_unicode(line) + r"\par" for line in lines)
    return (
        r"{\rtf1\ansi\deff0{\fonttbl{\f0 Times New Roman;}}{\stylesheet{\s0 Normal;}}"
        r"{\*\generator FixtureGen;}{\pict\pngblip 89504e470d0a1a0a0000PICTNOISE}"
        + "\n" + body + "\n}"
    ).encode("ascii")


def docx_bytes(title, paragraphs, table=None, merged=True, headings=()):
    import docx

    document = docx.Document()
    document.add_heading(title, level=1)
    for level, text in headings:
        document.add_heading(text, level=level)
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        rows, cols = len(table), len(table[0])
        tbl = document.add_table(rows=rows, cols=cols)
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                tbl.cell(r, c).text = value
        if merged and cols >= 2:
            merged_cell = tbl.cell(0, 0).merge(tbl.cell(0, 1))
            merged_cell.text = table[0][0]
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def xlsx_bytes(rows, sheet="Sheet1"):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def zip_bytes(members: dict[str, bytes]):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in members.items():
            zf.writestr(name, data)
    return buf.getvalue()


def soffice_path():
    for candidate in (os.environ.get("SEAL_SOFFICE"), "/Applications/LibreOffice.app/Contents/MacOS/soffice", shutil.which("soffice")):
        if candidate and Path(candidate).exists():
            return candidate
    return None


def soffice_convert(data: bytes, src_ext: str, target: str, workdir: Path) -> bytes:
    """Use LibreOffice (fixture generation only) to make .doc/.xls test inputs."""
    exe = soffice_path()
    if not exe:
        raise NotVerified("LibreOffice unavailable for fixture generation")
    workdir.mkdir(parents=True, exist_ok=True)
    tag = secrets.token_hex(4)
    src = workdir / f"in-{tag}.{src_ext}"
    src.write_bytes(data)
    profile = workdir / f"lo-profile-{tag}"
    subprocess.run(
        [exe, f"-env:UserInstallation=file://{profile}", "--headless", "--norestore", "--convert-to", target, "--outdir", str(workdir), str(src)],
        check=True, capture_output=True, timeout=180,
    )
    ext = target.split(":")[0]
    out = workdir / f"in-{tag}.{ext}"
    return out.read_bytes()


class PdfBuilder:
    """Tiny PDF writer: pages with Helvetica text lines and optional JPEG images."""

    def __init__(self):
        self.pages = []

    def page(self, lines=(), image_jpeg: bytes | None = None, image_size=(0, 0), text_size=12, small_text=None):
        self.pages.append({"lines": list(lines), "image": image_jpeg, "image_size": image_size, "size": text_size, "small": small_text})
        return self

    @staticmethod
    def _esc(text):
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    def build(self) -> bytes:
        objects: list[bytes] = []

        def add(obj: bytes) -> int:
            objects.append(obj)
            return len(objects)

        catalog = add(b"")  # placeholder 1
        pages_obj = add(b"")  # placeholder 2
        font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
        kids = []
        for page in self.pages:
            ops = []
            resources = f"/Font << /F1 {font} 0 R >>"
            if page["image"]:
                w, h = page["image_size"]
                img = add(
                    f"<< /Type /XObject /Subtype /Image /Width {w} /Height {h} /ColorSpace /DeviceGray "
                    f"/BitsPerComponent 8 /Filter /DCTDecode /Length {len(page['image'])} >>\nstream\n".encode()
                    + page["image"] + b"\nendstream"
                )
                resources += f" /XObject << /Im1 {img} 0 R >>"
                ops.append("q 612 0 0 792 0 0 cm /Im1 Do Q")
            y = 740
            for line in page["lines"]:
                ops.append(f"BT /F1 {page['size']} Tf 50 {y} Td ({self._esc(line)}) Tj ET")
                y -= page["size"] + 6
            if page["small"]:
                ops.append(f"BT /F1 4 Tf 500 20 Td ({self._esc(page['small'])}) Tj ET")
            stream = "\n".join(ops).encode("latin-1")
            content = add(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
            kids.append(add(
                f"<< /Type /Page /Parent {pages_obj} 0 R /MediaBox [0 0 612 792] /Resources << {resources} >> /Contents {content} 0 R >>".encode()
            ))
        objects[catalog - 1] = f"<< /Type /Catalog /Pages {pages_obj} 0 R >>".encode()
        objects[pages_obj - 1] = (
            "<< /Type /Pages /Kids [" + " ".join(f"{k} 0 R" for k in kids) + f"] /Count {len(kids)} >>"
        ).encode()
        raw, offsets = b"%PDF-1.4\n", []
        for index, obj in enumerate(objects, 1):
            offsets.append(len(raw))
            raw += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
        start = len(raw)
        raw += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
        raw += b"".join(f"{n:010d} 00000 n \n".encode() for n in offsets)
        raw += f"trailer\n<< /Size {len(objects) + 1} /Root {catalog} 0 R >>\nstartxref\n{start}\n%%EOF\n".encode()
        return raw


def _font(size):
    from PIL import ImageFont

    for path in (
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
        "/Library/Fonts/Arial.ttf",
    ):
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def scanned_page_jpeg(lines: list[str], width=1275, height=1650, font_size=44):
    """Render text as a grayscale 'scan' (150 DPI letter page)."""
    from PIL import Image, ImageDraw

    img = Image.new("L", (width, height), 255)
    draw = ImageDraw.Draw(img)
    font = _font(font_size)
    y = 120
    for line in lines:
        draw.text((100, y), line, fill=0, font=font)
        y += int(font_size * 1.6)
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=90)
    return buf.getvalue(), (width, height)


def noise_page_jpeg(width=1275, height=1650):
    """High-ink page with no recognisable text (hatching pattern)."""
    from PIL import Image, ImageDraw

    img = Image.new("L", (width, height), 255)
    draw = ImageDraw.Draw(img)
    for x in range(0, width + height, 9):
        draw.line([(x, 0), (x - height, height)], fill=30, width=3)
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=85)
    return buf.getvalue(), (width, height)


