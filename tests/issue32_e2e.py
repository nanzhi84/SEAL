"""Run: source /Users/yoryon/SEAL/work/dev/env.sh; PYTHONPATH=src /Users/yoryon/SEAL/.venv/bin/python tests/issue32_e2e.py"""
from __future__ import annotations
import json, re, shutil, sys
from pathlib import Path
from openpyxl import load_workbook
from issue32_support import HEADERS, Lab, sha
from issue32_faults import audit_manifest, concurrent_b, storage_faults, offline_faults, capacity, source_scopes, collision_subsets

lab=Lab()
try:
    lab.cli('init')
    lab.cli('collect',lab.config(),'--no-progress')
    state=lab.state()
    lab.check('real CLI collect produced three successful resources',state['tables']['version']['count']==3,state['tables']['version'])
    snap=lab.snapshot(); before=lab.state(); requests=len(lab.requests)
    package=lab.out/'delivery'
    lab.cli('export','--snapshot',snap,'--output',package)
    lab.check('export before publish is SQL/CAS read-only and never fetches',lab.state()==before and len(lab.requests)==requests)
    manifest=json.loads((package/'manifest.json').read_text())
    (lab.out/'manifest-observed.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False))
    wb=load_workbook(package/'metadata.xlsx',data_only=False); ws=wb.active
    lab.check('exact 16 headers',list(next(ws.values))==HEADERS)
    template=load_workbook(lab.work.parents[1]/'docs/reference/法规案例元数据上传模板.xlsx',read_only=True).active
    lab.check('template guidance row preserved exactly',list(ws.values)[1]==list(template.values)[1])
    lab.check('manifest data starts after preserved template guidance',manifest['data_start_row']==3)
    rows=[dict(zip(HEADERS,[v if v is not None else '' for v in row])) for row in list(ws.values)[2:]]
    lab.check('five delivery rows for single, two volumes and two independent files',len(rows)==5,[{'title':r['资料名称'],'order':r['文件顺序'],'file':r['解析后文件']} for r in rows])
    allowed={'法律法规','政策文件','案例','行业标准','通知公告','财务报告','会计准则','税务资料','其他'}
    for row_number,row in enumerate(rows,3):
        for key in HEADERS[:5]: lab.check(f'row {row_number} required {key}',bool(row[key]))
        lab.check(f'row {row_number} controlled type',row['资料类型'] in allowed)
        lab.check(f'row {row_number} collected format',bool(re.fullmatch(r'\d{4}-\d\d-\d\d \d\d:\d\d:\d\d',row['采集时间'])))
        lab.check(f'row {row_number} dates',all(not row[k] or re.fullmatch(r'\d{4}-\d\d-\d\d',row[k]) for k in ('发布日期','生效日期','失效日期')))
        lab.check(f'row {row_number} preserved text and no formula',row['文号/案号'] in ('','000042') and all(c.data_type!='f' for c in ws[row_number]))
        suffix=f"_{int(row['文件顺序']):02d}" if row['文件顺序'] else ''
        stem=f"{row['资料名称']}_{row['文件哈希'][:8]}{suffix}"
        lab.check(f'row {row_number} renamed filenames',row['解析后文件']==stem+'.md' and row['原始文件'].startswith(stem+'.'))
        md=(package/'markdown'/row['解析后文件']).read_bytes()
        original=(package/'originals'/row['原始文件']).read_bytes()
        lab.check(f'row {row_number} original SHA-256',sha(original)==row['文件哈希'])
    volumes=[r for r in rows if r['文件顺序']]
    lab.check('continuous volumes keep title, record and order 1..2',len(volumes)==2 and {int(r['文件顺序']) for r in volumes}=={1,2} and len({r['资料名称'] for r in volumes})==1 and len({r['来源记录号'] for r in volumes})==1)
    lab.check('independent attachment retains its own title and empty order',any(r['资料名称']=='独立修订说明' and not r['文件顺序'] for r in rows))
    audit_manifest(lab,package,manifest,rows,snap,before)
    lab.cli('export','verify',package)
    repeated=lab.cli('export','--snapshot',snap,'--output',package,ok=None)
    lab.check('same destination is idempotent',repeated.returncode==0,fatal=False)
    second=lab.out/'delivery-again'; lab.cli('export','--snapshot',snap,'--output',second)
    tree=lambda root:{str(p.relative_to(root)):sha(p.read_bytes()) for p in root.rglob('*') if p.is_file()}
    lab.check('same parameters produce byte-identical package',tree(package)==tree(second),fatal=False)
    lab.cli('snapshot','publish',snap); lab.api()
    response=json.loads(lab.get('/results?limit=100'))
    api_rows=response['items']; lab.check('API snapshot and row count match',response['snapshot_id']==snap and len(api_rows)==len(rows))
    byname={r['解析后文件']:r for r in api_rows}
    for row in rows:
        api=byname[row['解析后文件']]
        lab.check('API 16 columns exactly match workbook '+row['解析后文件'],all(str(api[k])==str(row[k]) for k in HEADERS))
        for folder,key,digest in [('markdown','解析后文件',api['markdown_hash']),('originals','原始文件',api['文件哈希'])]:
            lab.check('API file bytes match '+row[key],lab.get('/files/'+digest)==(package/folder/row[key]).read_bytes())
    moved=lab.out/'moved'/'delivery'; moved.parent.mkdir(); shutil.move(str(second),str(moved))
    network_guard=lab.work/'offline-network-guard'; network_guard.mkdir()
    (network_guard/'sitecustomize.py').write_text('import socket, os\nfrom pathlib import Path\nPath(os.environ["ACCEPT32_NET_ATTEMPTS"]+".loaded").write_text("network guard active")\ndef denied(*args, **kwargs):\n Path(os.environ["ACCEPT32_NET_ATTEMPTS"]).write_text("network attempt blocked")\n raise OSError("acceptance network disabled")\nsocket.socket.connect=denied\nsocket.socket.connect_ex=denied\nsocket.getaddrinfo=denied\n')
    offline=lab.env.copy(); offline['PYTHONPATH']=str(network_guard)+':'+offline['PYTHONPATH']; offline['ACCEPT32_NET_ATTEMPTS']=str(lab.work/'network-attempts.txt'); offline.update(SEAL_DATABASE_URL='postgresql://invalid:invalid@127.0.0.1:1/invalid',SEAL_S3_ENDPOINT='http://127.0.0.1:1',SEAL_S3_ACCESS_KEY='invalid',SEAL_S3_SECRET_KEY='invalid')
    lab.cli('export','verify',moved,env=offline,timeout=10)
    lab.check('moved package verifies with all socket connections disabled',(lab.work/'network-attempts.txt.loaded').exists() and not (lab.work/'network-attempts.txt').exists())
    corrupt=lab.out/'corrupt'; shutil.copytree(package,corrupt); first=next((corrupt/'markdown').iterdir()); first.write_bytes(first.read_bytes()+b'corrupted')
    lab.cli('export','verify',corrupt,ok=False); lab.check('offline verify refuses corrupted payload',True)
    lab.cli('export','--snapshot',snap,'--output',corrupt,ok=False); lab.check('existing different package is not overwritten',first.read_bytes().endswith(b'corrupted'))
    lab.cli('export','--snapshot',snap,'--source','dd-999','--output',lab.out/'wrong-scope',ok=False)
    lab.check('unknown source rejects without final directory',not (lab.out/'wrong-scope').exists())
    # Collect and publish B, then explicitly export A again: isolates published pointer.
    snap_b=concurrent_b(lab,snap,package)
    before=lab.state(); final_a=lab.out/'after-B'; lab.cli('export','--snapshot',snap,'--output',final_a)
    lab.check('A remains byte-identical after collect and publish B',tree(final_a)==tree(package),fatal=False)
    lab.check('A export preserves every SQL row and S3 object after B',lab.state()==before)
    lab.check('API still points to published B',json.loads(lab.get('/results?limit=100'))['snapshot_id']==snap_b)
    # Unknown volume order must fail during real collect, never create partial success.
    previous=lab.state()['tables']['version']['count']; lab.cli('collect',lab.config('/detail/ambiguous','ambiguous'),'--no-progress',ok=False)
    lab.check('ambiguous volume order creates no success version',lab.state()['tables']['version']['count']==previous)
    storage_faults(lab,snap,manifest)
    offline_faults(lab,package)
    capacity(lab)
    source_scopes(lab)
    collision_subsets(lab)
    # Package-only leak scan; do not put credential values into report.
    forbidden=[str(lab.work).encode(),str(Path.cwd()).encode(),lab.env['SEAL_S3_ACCESS_KEY'].encode(),lab.env['SEAL_S3_SECRET_KEY'].encode()]
    lab.check('delivery has no workspace path or private storage credentials',not any(value and value in p.read_bytes() for value in forbidden for p in package.rglob('*') if p.is_file()))
    (lab.out/'unverified.json').write_text(json.dumps({'real_export_scenarios_not_reached':['No independent original / no detail URL: deferred to issue33 real producer acceptance','Excel capacity: Source config rejects 32768-character title before collection; exporter capacity boundary not reached'],'not_claimed':'These are not counted as passed end-to-end exporter scenarios'},indent=2))
    (lab.out/'reproduce.txt').write_text('Prerequisites: Python 3.12 shared .venv, local PostgreSQL with CREATE DATABASE and isolated S3 bucket create/get/put/delete/list permissions.\nRequired variables: SEAL_DATABASE_URL, SEAL_S3_ENDPOINT, SEAL_S3_ACCESS_KEY, SEAL_S3_SECRET_KEY. Optional: SEAL_S3_REGION (default us-east-1).\ncd /Users/yoryon/SEAL-pr2\nsource /Users/yoryon/SEAL/work/dev/env.sh\nPYTHONPATH=src /Users/yoryon/SEAL/.venv/bin/python tests/issue32_e2e.py\nHarness creates fresh database/bucket, sets SEAL_ALLOW_LOOPBACK=1 and a run-owned SEAL_WORK_DIR. Fixtures and private API config stay in acceptance work/.\n')
    lab.finish()
    if any(not c['passed'] for c in lab.checks): sys.exit(1)
except Exception as exc:
    lab.finish(f'{type(exc).__name__}: {exc}')
    raise
