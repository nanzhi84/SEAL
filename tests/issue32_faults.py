"""External fault injection and independent manifest checks for issue32_e2e.py."""
import copy, http.client, json, os, shutil, signal, subprocess, threading, time, zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse
from openpyxl import load_workbook
import psycopg
from issue32_support import HEADERS, PYTHON, ROOT, port, sha


def audit_manifest(lab, package, manifest, rows, snap, state):
    lab.check('manifest schema, scope, counts and UNKNOWN coverage',manifest['schema']=='seal.delivery.v1' and manifest['snapshot_id']==snap and manifest['sources']==['dd-017'] and manifest['counts']=={'documents':3,'rows':5,'markdown_files':5,'original_files':5} and manifest['coverage']['dd-017']['site_population']=='UNKNOWN')
    actual={str(p.relative_to(package)):p for p in package.rglob('*') if p.is_file() and p.name!='manifest.json'}
    indexed={f['path']:f for f in manifest['files']}
    lab.check('manifest lists every payload exactly once',set(actual)==set(indexed) and len(indexed)==len(manifest['files']))
    for path,file in actual.items():
        item=indexed[path]; body=file.read_bytes()
        lab.check('manifest bytes and SHA-256 '+path,item['bytes']==len(body) and item['sha256']==sha(body))
    with psycopg.connect(lab.env['SEAL_DATABASE_URL']) as conn:
        conn.execute('SET TRANSACTION READ ONLY')
        sources=[r[0] for r in conn.execute('select row_to_json(t) from seal.snapshot_source t')]
        versions={str(r[0]):r[1] for r in conn.execute('select id, inputs_ref from seal.version')}
    lab.check('coverage is frozen snapshot coverage',all(manifest['coverage'][r['source_id']]==r['coverage'] for r in sources if str(r['snapshot_id'])==snap))
    for i,row in enumerate(rows,3):
        mapping=manifest['rows'][i-3]
        lab.check(f'manifest row {i} lineage fields',mapping['row']==i and mapping['source_id']=='dd-017' and mapping['source_record_hash']==row['来源记录号'] and mapping['file_hash']==row['文件哈希'] and mapping['markdown_file']=='markdown/'+row['解析后文件'] and mapping['original_file']=='originals/'+row['原始文件'] and mapping['markdown_hash']==sha((package/mapping['markdown_file']).read_bytes()))
        document=json.loads(lab.s3.get_object(Bucket=lab.bucket,Key=versions[mapping['version_id']])['Body'].read())
        def contains(value):
            if value==mapping['locator']: return True
            if isinstance(value,dict): return any(contains(v) for v in value.values())
            if isinstance(value,list): return any(contains(v) for v in value)
            return False
        lab.check(f'manifest row {i} locator exists in actual inputs_ref',contains(document))
    (lab.out/'readonly-before.json').write_text(json.dumps(state,indent=2,ensure_ascii=False))
    (lab.out/'readonly-after.json').write_text(json.dumps(lab.state(),indent=2,ensure_ascii=False))
    (lab.out/'hash-ledger.json').write_text(json.dumps(manifest['files'],indent=2,ensure_ascii=False))


class Gate:
    """Acceptance-owned network boundary: pause first S3 GET without patching seal."""
    def __init__(self,lab):
        self.entered=threading.Event(); self.release=threading.Event(); target=urlparse(lab.env['SEAL_S3_ENDPOINT']); gate=self
        class Proxy(BaseHTTPRequestHandler):
            def do_GET(self):
                gate.entered.set(); gate.release.wait(45)
                con=http.client.HTTPConnection(target.hostname,target.port or 80,timeout=20)
                try:
                    con.request('GET',self.path,headers=dict(self.headers)); res=con.getresponse(); data=res.read()
                    self.send_response(res.status)
                    for key,value in res.getheaders():
                        if key.lower() not in ('transfer-encoding','connection','date','server'): self.send_header(key,value)
                    self.end_headers(); self.wfile.write(data)
                except (BrokenPipeError,ConnectionResetError): pass
                finally: con.close()
            def log_message(self,*args): pass
        self.port=port(); self.http=ThreadingHTTPServer(('127.0.0.1',self.port),Proxy); self.http.daemon_threads=True
        threading.Thread(target=self.http.serve_forever,daemon=True).start()
        self.env=lab.env.copy(); self.env['SEAL_S3_ENDPOINT']=f'http://127.0.0.1:{self.port}'
    def close(self): self.release.set(); self.http.shutdown()


def concurrent_b(lab,snap,package):
    gate=Gate(lab); target=lab.out/'during-B'
    proc=subprocess.Popen([str(PYTHON),'-m','seal','export','--snapshot',snap,'--output',str(target)],cwd=ROOT,env=gate.env,text=True,stdout=subprocess.PIPE,stderr=subprocess.PIPE)
    try:
        lab.check('concurrent export reached controlled S3 boundary',gate.entered.wait(10))
        lab.cli('collect',lab.config('/detail/later','later'),'--no-progress'); snap_b=lab.snapshot(); lab.cli('snapshot','publish',snap_b)
        before=lab.state(); gate.release.set(); stdout,stderr=proc.communicate(timeout=30)
        lab.log.append({'args':['export','--snapshot',snap,'--output',str(target),'[S3 paused while B collected/published]'],'exit_code':proc.returncode,'stdout':stdout,'stderr':stderr})
        lab.check('concurrent A export succeeded',proc.returncode==0,{'exit_code':proc.returncode})
        lab.check('concurrent A export changed no SQL/CAS after B publish',lab.state()==before)
        a=load_workbook(package/'metadata.xlsx',data_only=False).active; b=load_workbook(target/'metadata.xlsx',data_only=False).active
        lab.check('concurrent export retains only A rows',list(a.values)==list(b.values))
        expected={str(p.relative_to(package)):sha(p.read_bytes()) for p in package.rglob('*') if p.is_file() and p.suffix in ('.md','.docx','.html')}
        actual={str(p.relative_to(target)):sha(p.read_bytes()) for p in target.rglob('*') if p.is_file() and p.suffix in ('.md','.docx','.html')}
        lab.check('concurrent A original and Markdown bytes are frozen',expected==actual)
        return snap_b
    finally:
        gate.close()
        if proc.poll() is None: proc.kill(); proc.wait()


def storage_faults(lab,snap,manifest):
    key=manifest['rows'][0]['file_hash']; obj=lab.s3.get_object(Bucket=lab.bucket,Key=key); body=obj['Body'].read(); media=obj['ContentType']
    try:
        lab.s3.put_object(Bucket=lab.bucket,Key=key,Body=body+b'CORRUPT',ContentType=media)
        target=lab.out/'bad-object'; receipt=lab.cli('export','--snapshot',snap,'--output',target,ok=False)
        lab.check('CAS hash corruption rejects atomically',not target.exists(),{'receipt':receipt.stdout})
        lab.s3.delete_object(Bucket=lab.bucket,Key=key)
        target=lab.out/'missing-object'; receipt=lab.cli('export','--snapshot',snap,'--output',target,ok=False)
        lab.check('CAS missing object rejects atomically',not target.exists(),{'receipt':receipt.stdout})
    finally: lab.s3.put_object(Bucket=lab.bucket,Key=key,Body=body,ContentType=media)
    gate=Gate(lab); target=lab.out/'cancelled'
    proc=subprocess.Popen([str(PYTHON),'-m','seal','export','--snapshot',snap,'--output',str(target)],cwd=ROOT,env=gate.env,text=True,stdout=subprocess.PIPE,stderr=subprocess.PIPE)
    try:
        lab.check('cancellable export reached S3 boundary',gate.entered.wait(10))
        proc.send_signal(signal.SIGINT); time.sleep(.2); gate.release.set(); stdout,stderr=proc.communicate(timeout=20)
        lab.check('cancelled export exits nonzero without final directory',proc.returncode!=0 and not target.exists(),{'exit_code':proc.returncode})
    finally:
        gate.close()
        if proc.poll() is None: proc.kill(); proc.wait()


def offline_faults(lab,package):
    def fixture(name):
        target=lab.out/name; shutil.copytree(package,target); return target
    target=fixture('missing-file'); next((target/'markdown').iterdir()).unlink()
    lab.cli('export','verify',target,ok=False); lab.check('verify refuses missing payload',True)
    target=fixture('extra-file'); (target/'extra.txt').write_text('not in manifest')
    lab.cli('export','verify',target,ok=False); lab.check('verify refuses unlisted payload',True)
    for name,col,value in [('invalid-enum',3,'INVALID_TYPE'),('formula-cell',9,'=1+1'),('bad-date',10,'2026-13-80'),('bad-time',5,'2026-10-20 25:61:99')]:
        target=fixture(name); book=load_workbook(target/'metadata.xlsx'); book.active.cell(3,col,value); book.save(target/'metadata.xlsx')
        manifest=json.loads((target/'manifest.json').read_text())
        for item in manifest['files']:
            if item['path']=='metadata.xlsx': item.update(bytes=(target/'metadata.xlsx').stat().st_size,sha256=sha((target/'metadata.xlsx').read_bytes()))
        (target/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False))
        receipt=lab.cli('export','verify',target,ok=None)
        lab.check('verify semantic refusal '+name,receipt.returncode!=0,{'receipt':receipt.stdout},fatal=False)
    target=fixture('bad-document-count'); manifest=json.loads((target/'manifest.json').read_text()); manifest['counts']['documents']+=1
    (target/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False))
    receipt=lab.cli('export','verify',target,ok=None)
    lab.check('verify refuses incorrect manifest document count',receipt.returncode!=0,{'receipt':receipt.stdout},fatal=False)
    target=fixture('unsafe-manifest'); manifest=json.loads((target/'manifest.json').read_text()); manifest['files'][0]['path']='../escape.md'
    (target/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False)); lab.cli('export','verify',target,ok=False)
    lab.check('verify refuses traversal manifest path',True)


def capacity(lab):
    lab.files['/detail/oversized']=lab.files['/detail/single']
    receipt=lab.cli('collect',lab.config('/detail/oversized','oversized',title='X'*32768),'--no-progress',ok=None)
    if receipt.returncode:
        lab.check('capacity fixture producer explicitly refused oversize before export',True,{'coverage':'export-capacity not reached','receipt':json.loads(receipt.stdout).get('errors')}); return
    snap=lab.snapshot(); target=lab.out/'capacity-exceeded'; receipt=lab.cli('export','--snapshot',snap,'--output',target,ok=None)
    lab.check('oversized Excel cell fails without final directory',receipt.returncode!=0 and not target.exists(),{'receipt':receipt.stdout},fatal=False)


def source_scopes(lab):
    lab.cli('collect',lab.config('/detail/single','second-source',title='../范围/隔离:资料*?',source='dd-040'),'--no-progress')
    snap=lab.snapshot(); target=lab.out/'source-subset'
    lab.cli('export','--snapshot',snap,'--source','dd-040','--output',target)
    manifest=json.loads((target/'manifest.json').read_text())
    lab.check('unsafe title is sanitized within package paths',all(not Path(item['path']).is_absolute() and '..' not in Path(item['path']).parts and (target/item['path']).resolve().is_relative_to(target.resolve()) for item in manifest['files']))
    lab.check('source subset is explicit and contains only selected source',manifest['sources']==['dd-040'] and manifest['counts']['rows']==1 and all(row['source_id']=='dd-040' for row in manifest['rows']))
    duplicate=lab.out/'source-repeat'
    lab.cli('export','--snapshot',snap,'--source','dd-040','--source','dd-040','--output',duplicate)
    repeat=json.loads((duplicate/'manifest.json').read_text())
    lab.check('repeated same source does not duplicate rows',repeat['sources']==['dd-040'] and repeat['counts']['rows']==1)
    multi=lab.out/'source-multi'
    lab.cli('export','--snapshot',snap,'--source','dd-040','--source','dd-017','--output',multi)
    manifest=json.loads((multi/'manifest.json').read_text())
    lab.check('multiple sources are sorted and complete',manifest['sources']==['dd-017','dd-040'] and {row['source_id'] for row in manifest['rows']}=={'dd-017','dd-040'})


def collision_subsets(lab):
    # A real second source deliberately delivers the same title and physical bytes.
    lab.cli('collect',lab.config('/detail/single','colliding-source',title='单件法规',source='dd-041'),'--no-progress')
    snap=lab.snapshot(); whole=lab.out/'collision-whole'
    lab.cli('export','--snapshot',snap,'--output',whole)
    full=json.loads((whole/'manifest.json').read_text())
    whole_rows=load_workbook(whole/'metadata.xlsx',read_only=True).active
    titles=[r[0] for r in list(whole_rows.values)[2:]]
    lab.check('collision fixture contains two distinct same-title records',titles.count('单件法规')==2)
    lab.check('whole snapshot names are batch-unique',len({row['markdown_file'] for row in full['rows']})==len(full['rows']) and len({row['original_file'] for row in full['rows']})==len(full['rows']))
    for source in ('dd-017','dd-041'):
        subset=lab.out/('collision-subset-'+source)
        lab.cli('export','--snapshot',snap,'--source',source,'--output',subset)
        partial=json.loads((subset/'manifest.json').read_text())
        identity=lambda row:(row['source_id'],row['resource_id'],row['version_id'],row['role'],row['volume_order'])
        expected={identity(row):(row['markdown_file'],row['original_file']) for row in full['rows'] if row['source_id']==source}
        actual={identity(row):(row['markdown_file'],row['original_file']) for row in partial['rows']}
        lab.check('source subset preserves whole-snapshot collision names '+source,expected==actual)
        lab.check('source subset preserves all collision payload bytes '+source,all((whole/row[k]).read_bytes()==(subset/row[k]).read_bytes() for row in partial['rows'] for k in ('markdown_file','original_file')))
