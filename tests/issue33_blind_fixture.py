"""Controlled complete legal-text sources for #33 public CLI acceptance."""
from __future__ import annotations
import io, json, threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlsplit, parse_qs
from docx import Document

UNICODE_HEADER = 'CHÍNH PHỦ — Hà Nội, ngày 20 tháng 10 năm 2026'
BODY_A = 'Controlled rule A. Article 1. This full legal text requires every operator to retain the original transaction evidence for seven years. Article 2. This rule enters into force on 20 October 2026. Article 3. The competent authority shall publish implementation guidance and supervise compliance. '
BODY_B = BODY_A.replace('rule A', 'rule B').replace('seven years','nine years')

def docx_bytes(text):
    d=Document();d.add_heading(UNICODE_HEADER,0);d.add_heading('Chương I',1);d.add_heading('Điều 1. Nội dung thực hiện',2);d.add_paragraph(text)
    stream=io.BytesIO();d.save(stream);return stream.getvalue()

class Fixture:
    def __init__(self):
        self.mode='normal';self.api_variant=0;self.events=[]
        self.files={
          '/files/part1.docx':docx_bytes('Các cơ quan có trách nhiệm thực hiện đầy đủ quy định của pháp luật. '+BODY_A),
          '/files/part2.docx':docx_bytes('Điều 2. Cơ quan ban hành chịu trách nhiệm công bố và tổ chức thực hiện văn bản. '+BODY_B),
          '/files/annex.rtf':b'{\\rtf1\\ansi{\\fonttbl{\\f0 Times New Roman;}}\\f0 Annex: The attached schedule contains the complete implementation duties, reporting periods, responsible authorities and the official record retention procedure.\\par Article 1. Every authority shall retain a complete implementation register.}',
          '/files/mixed.rtf':('{\\rtf1\\ansi\\ansicpg1252{\\fonttbl{\\f0 .VnTime;}}'+
             '\\u67?\\u72?\\u205?\\u78?\\u72? PH\\u7910? — Unicode header remains intact.\\par '+
             r"\'a7i\'d2u 1. Ban h\'b5nh quy \'ae\'denh v\'d2 quy\'d2n v\'b5 tr\'b8ch nhi\'d6m.\par "+
             'Article 2. The full legal rule requires each administrative authority to retain the original decision and to report implementation annually.}').encode('ascii',errors='replace'),
        }
        outer=self
        class Handler(BaseHTTPRequestHandler):
            def log_message(self,*args):pass
            def do_GET(self):outer.handle(self)
            def do_POST(self):outer.handle(self)
        self.server=ThreadingHTTPServer(('127.0.0.1',0),Handler)
        self.base=f'http://127.0.0.1:{self.server.server_port}'
        self.thread=threading.Thread(target=self.server.serve_forever,daemon=True);self.thread.start()
    def close(self):self.server.shutdown();self.server.server_close()
    def reply(self,h,data,kind='text/html; charset=utf-8',status=200):
        if isinstance(data,str):data=data.encode()
        h.send_response(status);h.send_header('Content-Type',kind);h.send_header('Content-Length',str(len(data)));h.end_headers();h.wfile.write(data)
    def handle(self,h):
        p=urlsplit(h.path);qs=parse_qs(p.query)
        data=h.rfile.read(int(h.headers.get('Content-Length','0'))).decode() if h.command=='POST' else ''
        self.events.append({'method':h.command,'path':h.path,'form':data,'mode':self.mode,'api_variant':self.api_variant})
        if p.path=='/robots.txt':return self.reply(h,'User-agent: *\nAllow: /','text/plain')
        if p.path=='/api/documents.json':
            a={'document_number':'2026-90001','title':'Controlled complete rule A','type':'Rule','agencies':[{'name':'Independent acceptance authority'}],'publication_date':'2026-10-20','body':BODY_A}
            b={'document_number':'2026-90002','title':'Controlled complete rule B','type':'Rule','agencies':[{'name':'Independent acceptance authority'}],'publication_date':'2026-10-20','body':BODY_B + ('Updated substantive Article 4. A report is required monthly.' if self.api_variant else '')}
            if self.mode=='api-quality':a['body']='Corrupt full legal text: '+('\ufffd'*40)
            if self.mode=='api-missing-body':a.pop('body')
            if self.mode=='api-structure':return self.reply(h,json.dumps({'summaries':[a,b]}),'application/json')
            items=[b,a] if self.api_variant else [a,b]
            return self.reply(h,json.dumps({'count':4 if self.mode=='api-pagination' else 2,'results':items,**({'next_page_url':self.base+'/api/documents.json?page=2'} if self.mode=='api-pagination' else {})}), 'application/json')
        if p.path.startswith('/files/'):
            if self.mode=='missing' and p.path=='/files/part2.docx':return self.reply(h,'missing','text/plain',404)
            if self.mode=='quality' and p.path=='/files/part2.docx':return self.reply(h,docx_bytes('Corrupted required legal text: '+('\ufffd'*40)),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            if p.path in self.files:return self.reply(h,self.files[p.path], 'application/octet-stream')
        if p.path=='/vn/':
            if 'docid' not in qs:
                if self.mode=='structure':return self.reply(h,'<html><h1>Legal portal redesigned</h1></html>')
                page=2 if h.command=='POST' else 1
                docid=9001 if page==1 else 9002
                nextlink='<a href="javascript:__doPostBack(\'grid\',\'Page$2\')">2</a>' if page==1 else ''
                return self.reply(h,f'<html><form method="post" action="/vn/"><input type="hidden" name="__VIEWSTATE" value="controlled-state"><table id="grvDocument"><tr><td><a href="?docid={docid}">Complete document</a></td></tr><tr><td><span>{page}</span>{nextlink}</td></tr></table></form></html>')
            number=qs['docid'][0]
            if self.mode=='challenge':return self.reply(h,'<html><title>Just a moment...</title><p>Verify you are human</p></html>')
            if number=='9001':
                second='Part 3' if self.mode=='ambiguous' else 'Part 2'
                href='https://example.com/part2.docx' if self.mode=='out-of-scope' else '/files/part2.docx'
                links=f'<a href="/files/part1.docx">Part 1</a><a href="{href}">{second}</a><a href="/files/annex.rtf">Independent annex</a>'
            else:links='<a href="/files/mixed.rtf">Historical font regulation</a>'
            return self.reply(h,f'<html><title>Quy định {number}</title><h4 class="title">Quy định {number}</h4><table><tr><td>Số ký hiệu</td><td>{number}/2026/NĐ-CP</td></tr><tr><td>Ngày ban hành</td><td>20-10-2026</td></tr><tr><td>Cơ quan ban hành</td><td>Chính phủ</td></tr><tr><td>Trích yếu</td><td>Metadata only, never the full legal body.</td></tr></table>{links}</html>')
        return self.reply(h,'not found','text/plain',404)
